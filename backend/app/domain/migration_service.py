"""
Phase 19: 기존 차입금 일괄 이관(관리자).

흐름: 양식(Excel) 내려받기 → 기존 자료 붙여넣기 → 검사(실제 등록·원리금 생성까지 돌려 본 뒤 되돌림) → 이관 실행(전부 아니면 전무)
→ 이관 대조표(Excel) → 필요하면 이관 취소(숨김 삭제, 삭제·복구에서 되살림).

- 기숙사/SPC: 같은 이름이 이미 있으면 그것을 쓰고, 없으면 만든다.
- 차입금: 기본조건 + 최초 금리 + (추가 금리이력) → 원리금 테이블 V1을 금리이력 전체로 생성(금리변동 이벤트 없이 처음부터 반영).
- 납부: '이관 기준일' 이전(포함) 상환예정일은 예정액대로 납부된 것으로 기록(선택). 예외 시트에 적은 회차는 그 금액으로(0이면 미납).
- 중도변수(상환유예·조기상환·가속상환·대환)는 건마다 근거를 확인해야 하므로 일괄 대상이 아니다 — 이관 뒤 이벤트 화면에서 순서대로 등록.
"""
import io
import re
import uuid
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.rate_service import RateNotRegisteredError
from app.domain.schedule_amounts import scheduled_by_due_date
from app.domain.schedule_service import ScheduleGenerationError, generate_initial_schedule
from app.domain.version_utils import get_current_version
from app.models.loan import InterestRateHistory, Loan, RateType
from app.models.organization import SPC, Dormitory, MigrationBatch, User
from app.models.payment import ActualPayment
from app.models.reference import Code, Fund

# ---------- 양식 ----------
SHEETS = {
    "기숙사": [("기숙사명*", "예) 가람관"), ("학교명*", "예) 가람대학교"), ("SPC명*", "예) 가람관(주) — 같은 이름이 있으면 그 SPC를 씀"),
             ("사업유형", "예) 행복기숙사(사립)"), ("담당자 이메일", "재단 담당자 로그인 아이디(선택)")],
    "차입금": [("기숙사명*", "기숙사 시트 또는 이미 등록된 기숙사 이름"), ("학교명", "같은 이름의 기숙사가 여럿일 때 구분"),
             ("차입금코드*", "예) 2017-08-21_공자-00 (중복 불가)"), ("재원*", "목록에서 선택"), ("배정금액(원)*", "숫자"),
             ("배정일*", "YYYY-MM-DD"), ("금리유형*", "고정 / 변동"), ("상환주기*", "목록에서 선택"), ("상환방식*", "목록에서 선택"),
             ("거치유형", "거치가 있으면 선택"), ("거치시작일", "YYYY-MM-DD"), ("거치종료일", "YYYY-MM-DD (마지막 거치 이자 납입일 권장)"),
             ("원금상환시작일*", "첫 원금 상환예정일"), ("상환기간(개월)*", "예) 324"), ("이자끝수처리", "반올림(기본) / 절사"),
             ("최초금리(%)*", "예) 2.21"), ("최초금리 적용시작일", "비우면 배정일")],
    "금리이력": [("차입금코드*", "변동금리 등 최초금리 이후 바뀐 금리만"), ("적용시작일*", "YYYY-MM-DD"), ("연금리(%)*", "예) 2.98")],
    "납부예외": [("차입금코드*", "기준일 이전 회차 중 예정대로 납부되지 않은 회차"), ("상환예정일*", "원리금 테이블의 상환예정일"),
             ("실제납부 원금(원)*", "미납이면 0"), ("실제납부 이자(원)*", "미납이면 0"), ("실제납부일", "비우면 상환예정일")],
}
GUIDE = [
    "기존 차입금 일괄 이관 양식",
    "",
    "1. '기숙사' → '차입금' → (변동금리면) '금리이력' 순서로 채웁니다. 별표(*) 칸은 꼭 채웁니다. 첫 줄(제목)은 지우지 마세요.",
    "2. 날짜는 YYYY-MM-DD(예: 2017-08-21), 금액은 숫자만(쉼표 가능), 금리는 % 숫자(예: 2.21)로 적습니다.",
    "3. 재원·상환주기·상환방식·거치유형·금리유형은 칸을 누르면 나오는 목록에서 고릅니다(시스템 기준정보와 같은 이름).",
    "4. 금리이력: 최초금리 이후에 바뀐 금리만 적습니다. 원리금 테이블은 이 이력을 모두 반영해 처음부터 만들어집니다.",
    "5. 납부: 화면에서 '이관 기준일'과 '기준일 이전 회차 납부완료 처리'를 고르면, 기준일까지의 회차는 예정액대로 납부된 것으로 기록됩니다.",
    "   예정대로 내지 않은 회차만 '납부예외' 시트에 적습니다(미납이면 원금·이자 0, 일부납부면 실제 금액).",
    "6. 상환유예·조기상환·가속상환·대환은 일괄 대상이 아닙니다. 이관 뒤 '이벤트 관리'에서 발생 순서대로 등록하세요.",
    "7. 화면에서 '검사'를 누르면 실제로 등록·계산까지 해 본 뒤 되돌립니다. 오류가 없을 때만 '이관 실행'이 됩니다(전부 반영 또는 전혀 반영 안 함).",
    "8. 이관 후 '이관 대조표'를 내려받아 기존 엑셀의 잔여원금·다음 상환액과 맞춰 보세요.",
]
HEAD_FILL = PatternFill("solid", fgColor="0F2A4D")


def build_template(db: Session) -> bytes:
    wb = Workbook()
    g = wb.active
    g.title = "안내"
    for i, line in enumerate(GUIDE, start=1):
        c = g.cell(row=i, column=1, value=line)
        if i == 1:
            c.font = Font(bold=True, size=14)
    g.column_dimensions["A"].width = 130

    lists = wb.create_sheet("목록")
    funds = [f.fund_name for f in db.execute(select(Fund).where(Fund.is_active.is_(True), Fund.is_deleted.is_(False))
                                             .order_by(Fund.fund_name)).scalars()]
    groups = {}
    for grp in ("REPAY_CYCLE", "REPAY_METHOD", "GRACE_TYPE"):
        groups[grp] = [c.label for c in db.execute(select(Code).where(Code.group_code == grp, Code.is_active.is_(True),
                                                                      Code.is_deleted.is_(False)).order_by(Code.sort_order)).scalars()]
    cols = [("재원", funds), ("금리유형", ["고정", "변동"]), ("상환주기", groups["REPAY_CYCLE"]),
            ("상환방식", groups["REPAY_METHOD"]), ("거치유형", groups["GRACE_TYPE"]), ("이자끝수처리", ["반올림", "절사"])]
    ranges = {}
    for ci, (name, vals) in enumerate(cols, start=1):
        lists.cell(row=1, column=ci, value=name).font = Font(bold=True)
        for ri, v in enumerate(vals, start=2):
            lists.cell(row=ri, column=ci, value=v)
        letter = lists.cell(row=1, column=ci).column_letter
        ranges[name] = f"목록!${letter}$2:${letter}${max(2, len(vals) + 1)}"
    lists.sheet_state = "hidden"

    for sheet, heads in SHEETS.items():
        ws = wb.create_sheet(sheet)
        for ci, (h, hint) in enumerate(heads, start=1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center")
            c.comment = Comment(hint, "이관 양식")
            ws.column_dimensions[c.column_letter].width = max(14, len(h) * 2 + 4)
            key = h.rstrip("*")
            if key in ranges:
                dv = DataValidation(type="list", formula1=f"={ranges[key]}", allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{c.column_letter}2:{c.column_letter}2000")
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- 읽기 ----------
class Issues:
    def __init__(self):
        self.errors: list[dict] = []
        self.warnings: list[dict] = []
        self.rows: dict[str, int] = {}

    def err(self, sheet, row, col, msg):
        self.errors.append({"sheet": sheet, "row": row, "column": col, "message": msg})

    def warn(self, sheet, row, col, msg):
        self.warnings.append({"sheet": sheet, "row": row, "column": col, "message": msg})


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _date(v):
    if v is None or _text(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    m = re.fullmatch(r"(\d{4})[-./ ]?(\d{1,2})[-./ ]?(\d{1,2})\.?", s)
    if not m:
        raise ValueError(f"날짜 형식이 아닙니다: '{s}' (YYYY-MM-DD)")
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        raise ValueError(f"없는 날짜입니다: '{s}'")


def _int(v):
    s = _text(v).replace(",", "").replace("원", "")
    if s == "":
        return None
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"숫자가 아닙니다: '{_text(v)}'")
    if d != d.to_integral_value():
        raise ValueError(f"원 단위 정수여야 합니다: '{_text(v)}'")
    return int(d)


def _rate(v):
    s = _text(v).replace("%", "")
    if s == "":
        return None
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"금리가 숫자가 아닙니다: '{_text(v)}'")
    if isinstance(v, float) and 0 < v < 0.2 and "%" not in _text(v):
        d = d * 100   # Excel 백분율 서식(0.0221)으로 들어온 값
    if not (Decimal(0) < d < Decimal(100)):
        raise ValueError(f"금리는 0~100 사이의 % 값이어야 합니다: '{_text(v)}'")
    return d.quantize(Decimal("0.00001"))


def _rows(wb, sheet, issues: Issues) -> list[tuple[int, dict]]:
    if sheet not in wb.sheetnames:
        if sheet in ("기숙사", "차입금"):
            issues.err(sheet, None, None, f"'{sheet}' 시트가 없습니다. 이관 양식을 내려받아 쓰세요.")
        return []
    ws = wb[sheet]
    heads = [_text(c.value).rstrip("*").strip() for c in ws[1]]
    want = [h.rstrip("*") for h, _ in SHEETS[sheet]]
    missing = [h for h in want if h not in heads]
    if missing:
        issues.err(sheet, 1, None, f"제목 줄에 없는 칸: {', '.join(missing)} — 양식의 제목을 바꾸지 마세요.")
        return []
    idx = {h: heads.index(h) for h in want}
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = {h: (row[i] if i < len(row) else None) for h, i in idx.items()}
        if all(_text(v) == "" for v in vals.values()):
            continue
        out.append((r, vals))
    issues.rows[sheet] = len(out)
    return out


def _need(issues, sheet, r, vals, *cols) -> bool:
    ok = True
    for c in cols:
        if _text(vals.get(c)) == "":
            issues.err(sheet, r, c, "필수 칸이 비어 있습니다.")
            ok = False
    return ok


def _parse(db: Session, content: bytes, issues: Issues) -> dict:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:  # noqa: BLE001
        issues.err("파일", None, None, "Excel(.xlsx) 파일을 읽지 못했습니다. 이관 양식을 내려받아 저장한 .xlsx 파일을 올리세요.")
        return {}
    funds = {f.fund_name: f for f in db.execute(select(Fund).where(Fund.is_active.is_(True), Fund.is_deleted.is_(False))).scalars()}
    codes = defaultdict(dict)
    for c in db.execute(select(Code).where(Code.is_active.is_(True), Code.is_deleted.is_(False))).scalars():
        codes[c.group_code][c.label] = c
        codes[c.group_code][c.code] = c

    dorms, loans, rates, excepts = [], [], [], []
    seen_dorm = set()
    for r, v in _rows(wb, "기숙사", issues):
        if not _need(issues, "기숙사", r, v, "기숙사명", "학교명", "SPC명"):
            continue
        key = (_text(v["기숙사명"]), _text(v["학교명"]))
        if key in seen_dorm:
            issues.err("기숙사", r, "기숙사명", f"같은 기숙사가 두 번 적혀 있습니다: {key[0]} ({key[1]})")
            continue
        seen_dorm.add(key)
        email = _text(v["담당자 이메일"]) or None
        mgr = None
        if email:
            mgr = db.execute(select(User).where(User.email == email, User.is_deleted.is_(False))).scalar_one_or_none()
            if mgr is None:
                issues.err("기숙사", r, "담당자 이메일", f"등록된 사용자가 아닙니다: {email}")
        dorms.append({"row": r, "name": key[0], "school": key[1], "spc": _text(v["SPC명"]),
                      "project_type": _text(v["사업유형"]) or None, "manager": mgr})

    seen_code = set()
    for r, v in _rows(wb, "차입금", issues):
        s = "차입금"
        if not _need(issues, s, r, v, "기숙사명", "차입금코드", "재원", "배정금액(원)", "배정일", "금리유형", "상환주기",
                     "상환방식", "원금상환시작일", "상환기간(개월)", "최초금리(%)"):
            continue
        item = {"row": r, "dorm": _text(v["기숙사명"]), "school": _text(v["학교명"]) or None, "code": _text(v["차입금코드"])}
        ok = True
        if len(item["code"]) > 30:
            issues.err(s, r, "차입금코드", "30자 이내여야 합니다.")
            ok = False
        if item["code"] in seen_code:
            issues.err(s, r, "차입금코드", f"같은 차입금코드가 두 번 적혀 있습니다: {item['code']}")
            ok = False
        seen_code.add(item["code"])
        if db.execute(select(Loan.id).where(Loan.loan_code == item["code"])).first():
            issues.err(s, r, "차입금코드", f"이미 등록된 차입금코드입니다: {item['code']}")
            ok = False
        fund = funds.get(_text(v["재원"]))
        if fund is None:
            issues.err(s, r, "재원", f"기준정보에 없는 재원입니다: '{_text(v['재원'])}'")
            ok = False
        rt = {"고정": RateType.FIXED, "고정금리": RateType.FIXED, "FIXED": RateType.FIXED,
              "변동": RateType.VARIABLE, "변동금리": RateType.VARIABLE, "VARIABLE": RateType.VARIABLE}.get(_text(v["금리유형"]))
        if rt is None:
            issues.err(s, r, "금리유형", "고정 또는 변동으로 적으세요.")
            ok = False
        cyc = codes["REPAY_CYCLE"].get(_text(v["상환주기"]))
        mth = codes["REPAY_METHOD"].get(_text(v["상환방식"]))
        if cyc is None:
            issues.err(s, r, "상환주기", f"기준정보에 없는 상환주기입니다: '{_text(v['상환주기'])}'")
            ok = False
        if mth is None:
            issues.err(s, r, "상환방식", f"기준정보에 없는 상환방식입니다: '{_text(v['상환방식'])}'")
            ok = False
        gt = None
        if _text(v["거치유형"]):
            gt = codes["GRACE_TYPE"].get(_text(v["거치유형"]))
            if gt is None:
                issues.err(s, r, "거치유형", f"기준정보에 없는 거치유형입니다: '{_text(v['거치유형'])}'")
                ok = False
        rounding = {"": "ROUND", "반올림": "ROUND", "ROUND": "ROUND", "절사": "ROUNDDOWN", "ROUNDDOWN": "ROUNDDOWN"}.get(_text(v["이자끝수처리"]))
        if rounding is None:
            issues.err(s, r, "이자끝수처리", "반올림 또는 절사로 적으세요.")
            ok = False
        vals = {}
        for col, fn, key in (("배정금액(원)", _int, "amount"), ("배정일", _date, "assigned"), ("거치시작일", _date, "gs"),
                             ("거치종료일", _date, "ge"), ("원금상환시작일", _date, "first"), ("상환기간(개월)", _int, "months"),
                             ("최초금리(%)", _rate, "rate"), ("최초금리 적용시작일", _date, "rate_start")):
            try:
                vals[key] = fn(v[col])
            except ValueError as e:
                issues.err(s, r, col, str(e))
                ok = False
        if ok:
            if vals["amount"] is None or vals["amount"] <= 0:
                issues.err(s, r, "배정금액(원)", "0보다 커야 합니다.")
                ok = False
            if vals["months"] is None or vals["months"] <= 0:
                issues.err(s, r, "상환기간(개월)", "0보다 커야 합니다.")
                ok = False
            if (vals["gs"] is None) != (vals["ge"] is None):
                issues.err(s, r, "거치시작일", "거치시작일과 거치종료일은 둘 다 적거나 둘 다 비워야 합니다.")
                ok = False
            elif vals["gs"] and vals["gs"] > vals["ge"]:
                issues.err(s, r, "거치종료일", "거치종료일이 거치시작일보다 빠릅니다.")
                ok = False
            if vals["first"] and vals["assigned"] and vals["first"] <= vals["assigned"]:
                issues.err(s, r, "원금상환시작일", "원금상환시작일은 배정일보다 뒤여야 합니다.")
                ok = False
            rs = vals["rate_start"] or vals["assigned"]
            if rs and vals["assigned"] and rs > vals["assigned"]:
                issues.warn(s, r, "최초금리 적용시작일", "배정일보다 늦습니다 — 첫 이자기간에 쓸 금리가 없으면 검사에서 오류가 납니다.")
        if ok:
            item.update(fund=fund, rate_type=rt, cycle=cyc, method=mth, grace=gt, rounding=rounding, **vals)
            item["rate_start"] = vals["rate_start"] or vals["assigned"]
            loans.append(item)

    loan_codes = {x["code"] for x in loans} | seen_code
    for r, v in _rows(wb, "금리이력", issues):
        if not _need(issues, "금리이력", r, v, "차입금코드", "적용시작일", "연금리(%)"):
            continue
        code = _text(v["차입금코드"])
        if code not in loan_codes:
            issues.err("금리이력", r, "차입금코드", f"'차입금' 시트에 없는 차입금코드입니다: {code}")
            continue
        try:
            rates.append({"row": r, "code": code, "start": _date(v["적용시작일"]), "rate": _rate(v["연금리(%)"])})
        except ValueError as e:
            issues.err("금리이력", r, None, str(e))

    for r, v in _rows(wb, "납부예외", issues):
        if not _need(issues, "납부예외", r, v, "차입금코드", "상환예정일", "실제납부 원금(원)", "실제납부 이자(원)"):
            continue
        code = _text(v["차입금코드"])
        if code not in loan_codes:
            issues.err("납부예외", r, "차입금코드", f"'차입금' 시트에 없는 차입금코드입니다: {code}")
            continue
        try:
            item = {"row": r, "code": code, "due": _date(v["상환예정일"]), "principal": _int(v["실제납부 원금(원)"]),
                    "interest": _int(v["실제납부 이자(원)"]), "paid": _date(v["실제납부일"])}
        except ValueError as e:
            issues.err("납부예외", r, None, str(e))
            continue
        if item["principal"] < 0 or item["interest"] < 0:
            issues.err("납부예외", r, None, "금액은 음수일 수 없습니다.")
            continue
        if any(x["code"] == code and x["due"] == item["due"] for x in excepts):
            issues.err("납부예외", r, "상환예정일", "같은 차입금·상환예정일이 두 번 적혀 있습니다.")
            continue
        excepts.append(item)
    return {"dorms": dorms, "loans": loans, "rates": rates, "excepts": excepts}


# ---------- 반영 ----------
@contextmanager
def _hold_commits(db: Session):
    """기존 업무 함수(원리금 생성 등)가 중간에 commit하지 않게 — 모두 끝난 뒤 한 번에 반영하거나 되돌린다."""
    db.commit = db.flush   # type: ignore[method-assign]
    try:
        yield
    finally:
        del db.commit


def _run(db: Session, data: dict, *, cutoff: date, mark_paid: bool, user: User, issues: Issues) -> dict:
    stats = {"spc_new": 0, "dorm_new": 0, "dorm_reused": 0, "loans": 0, "rates": 0, "versions": 0,
             "payments_full": 0, "payments_exception": 0, "unpaid_exceptions": 0}
    created = {"loans": [], "dormitories": []}
    spcs = {s.spc_name: s for s in db.execute(select(SPC).where(SPC.is_deleted.is_(False))).scalars()}
    existing = list(db.execute(select(Dormitory).where(Dormitory.is_deleted.is_(False))).scalars())
    dorm_by = {}
    for d in existing:
        dorm_by.setdefault((d.dormitory_name, d.school_name), d)
    for d in data["dorms"]:
        key = (d["name"], d["school"])
        if key in dorm_by:
            stats["dorm_reused"] += 1
            issues.warn("기숙사", d["row"], "기숙사명", f"이미 등록된 기숙사라 그대로 씁니다: {d['name']} ({d['school']})")
            continue
        spc = spcs.get(d["spc"])
        if spc is None:
            spc = SPC(spc_name=d["spc"], created_by=user.id)
            db.add(spc)
            db.flush()
            spcs[d["spc"]] = spc
            stats["spc_new"] += 1
        dm = Dormitory(dormitory_name=d["name"], school_name=d["school"], spc_id=spc.id, project_type=d["project_type"],
                       manager_user_id=d["manager"].id if d["manager"] else None, created_by=user.id)
        db.add(dm)
        db.flush()
        dorm_by[key] = dm
        created["dormitories"].append(str(dm.id))
        stats["dorm_new"] += 1

    rates_by = defaultdict(list)
    for x in data["rates"]:
        rates_by[x["code"]].append(x)
    exc_by = defaultdict(dict)
    for x in data["excepts"]:
        exc_by[x["code"]][x["due"]] = x
    for it in data["loans"]:
        cands = [d for (n, sc), d in dorm_by.items() if n == it["dorm"] and (it["school"] is None or sc == it["school"])]
        if not cands:
            issues.err("차입금", it["row"], "기숙사명", f"기숙사 시트에도, 등록된 기숙사에도 없는 이름입니다: {it['dorm']}")
            continue
        if len(cands) > 1:
            issues.err("차입금", it["row"], "학교명", f"'{it['dorm']}' 기숙사가 여럿입니다. 학교명을 적어 구분하세요.")
            continue
        loan = Loan(dormitory_id=cands[0].id, fund_id=it["fund"].id, loan_code=it["code"], assigned_amount=it["amount"],
                    assigned_date=it["assigned"], rate_type=it["rate_type"], repay_cycle_code_id=it["cycle"].id,
                    repay_method_code_id=it["method"].id, grace_type_code_id=it["grace"].id if it["grace"] else None,
                    grace_start_date=it["gs"], grace_end_date=it["ge"], principal_repay_start_date=it["first"],
                    repay_period_months=it["months"], interest_rounding=it["rounding"], created_by=user.id)
        db.add(loan)
        db.flush()
        hist = [{"start": it["rate_start"], "rate": it["rate"], "row": it["row"], "sheet": "차입금"}]
        for x in sorted(rates_by.get(it["code"], []), key=lambda x: x["start"]):
            if x["start"] <= hist[-1]["start"]:
                issues.err("금리이력", x["row"], "적용시작일", f"{it['code']}: 앞선 금리({hist[-1]['start']})보다 늦은 날짜여야 합니다.")
                continue
            hist.append({"start": x["start"], "rate": x["rate"], "row": x["row"], "sheet": "금리이력"})
        for i, h in enumerate(hist):
            end = hist[i + 1]["start"] if i + 1 < len(hist) else None
            db.add(InterestRateHistory(loan_id=loan.id, effective_start_date=h["start"], effective_end_date=end,
                                       annual_rate=h["rate"], created_by=user.id))
        db.flush()
        stats["rates"] += len(hist) - 1
        sp = db.begin_nested()
        try:
            generate_initial_schedule(db, loan.id, generated_by=user.id)
            sp.commit()
        except (ScheduleGenerationError, RateNotRegisteredError, ValueError, ArithmeticError) as e:
            sp.rollback()
            issues.err("차입금", it["row"], None, f"{it['code']}: 원리금 테이블을 만들지 못했습니다 — {e}")
            continue
        stats["loans"] += 1
        stats["versions"] += 1
        created["loans"].append(str(loan.id))

        sched = scheduled_by_due_date(db, get_current_version(db, loan.id))
        exc = exc_by.get(it["code"], {})
        for d, x in exc.items():
            if d not in sched or sched[d]["total"] <= 0:
                near = min((k for k in sched if sched[k]["total"] > 0), key=lambda k: abs((k - d).days), default=None)
                hint = f" 가까운 상환예정일: {near}" if near else ""
                issues.err("납부예외", x["row"], "상환예정일", f"{it['code']}: {d}은(는) 원리금 테이블의 상환예정일이 아닙니다.{hint}")
            elif d > cutoff:
                issues.err("납부예외", x["row"], "상환예정일", f"{it['code']}: 이관 기준일({cutoff}) 뒤의 회차입니다.")
        for d in sorted(sched):
            if d > cutoff or sched[d]["total"] <= 0:
                continue
            x = exc.get(d)
            if x is not None:
                if x["principal"] + x["interest"] == 0:
                    stats["unpaid_exceptions"] += 1
                    continue
                db.add(ActualPayment(loan_id=loan.id, due_date=d, actual_paid_date=x["paid"] or d,
                                     actual_principal=x["principal"], actual_interest=x["interest"],
                                     remark="기존 차입금 이관(예외 시트)", created_by=user.id))
                stats["payments_exception"] += 1
            elif mark_paid:
                db.add(ActualPayment(loan_id=loan.id, due_date=d, actual_paid_date=d, actual_principal=sched[d]["principal"],
                                     actual_interest=sched[d]["interest"], remark=f"기존 차입금 이관(기준일 {cutoff} 이전 예정대로 납부)",
                                     created_by=user.id))
                stats["payments_full"] += 1
        db.flush()
    return {"stats": stats, "created": created}


def run_migration(db: Session, content: bytes, file_name: str, *, cutoff: date, mark_paid: bool, user: User, apply: bool) -> dict:
    issues = Issues()
    if cutoff > date.today():
        issues.err("화면", None, "이관 기준일", "이관 기준일은 오늘 이후일 수 없습니다(미래 회차를 납부완료로 기록하게 됨).")
    data = _parse(db, content, issues)
    result = {"stats": {}, "created": {"loans": [], "dormitories": []}}
    if data and not issues.errors:
        if not data["loans"]:
            issues.err("차입금", None, None, "차입금 시트에 이관할 행이 없습니다.")
        else:
            with _hold_commits(db):
                try:
                    result = _run(db, data, cutoff=cutoff, mark_paid=mark_paid, user=user, issues=issues)
                except Exception:
                    db.rollback()
                    raise
    counts = {"dorm_rows": issues.rows.get("기숙사", 0), "loan_rows": issues.rows.get("차입금", 0),
              "rate_rows": issues.rows.get("금리이력", 0), "exception_rows": issues.rows.get("납부예외", 0)}
    out = {"ok": not issues.errors, "applied": False, "errors": issues.errors, "warnings": issues.warnings,
           "counts": counts, "stats": result["stats"], "batch_id": None}
    if apply and not issues.errors:
        batch = MigrationBatch(file_name=file_name[:200], cutoff_date=cutoff, mark_paid=mark_paid,
                               summary={**counts, **result["stats"]}, loan_ids=result["created"]["loans"],
                               dormitory_ids=result["created"]["dormitories"], created_by=user.id)
        db.add(batch)
        db.commit()
        out.update(applied=True, batch_id=str(batch.id))
    else:
        db.rollback()        # 검사는 항상, 실행은 오류가 있으면 아무것도 남기지 않는다
    return out


# ---------- 대조표 ----------
def build_report(db: Session, batch: MigrationBatch) -> bytes:
    from app.domain.verification_service import verify_loan
    wb = Workbook()
    ws = wb.active
    ws.title = "이관 대조표"
    head = ["차입금코드", "기숙사", "학교", "재원", "배정금액", "원리금 버전", "총 원금", "총 이자",
            f"잔여원금({batch.cutoff_date} 기준·예정)", "다음 상환예정일", "다음 회차 원리금", "기준일까지 회차", "납부 기록",
            "미납·부족 회차", "계산검증", "기존 엑셀 잔여원금(직접 입력)", "차이"]
    for ci, h in enumerate(head, start=1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r = 2
    for lid in batch.loan_ids or []:
        loan = db.get(Loan, uuid.UUID(lid))
        if loan is None or loan.is_deleted:
            continue
        dorm = db.get(Dormitory, loan.dormitory_id)
        fund = db.get(Fund, loan.fund_id)
        try:
            v = get_current_version(db, loan.id)
        except Exception:  # noqa: BLE001
            continue
        sched = scheduled_by_due_date(db, v)
        paid = defaultdict(int)
        n_pay = 0
        for p in db.execute(select(ActualPayment).where(ActualPayment.loan_id == loan.id,
                                                        ActualPayment.status == "ACTIVE")).scalars():
            paid[p.due_date] += p.actual_principal + p.actual_interest
            n_pay += 1
        past = [d for d in sched if d <= batch.cutoff_date and sched[d]["total"] > 0]
        short = sum(1 for d in past if paid[d] < sched[d]["total"])
        remain = loan.assigned_amount - sum(sched[d]["principal"] for d in past)
        nxt = min((d for d in sched if d > batch.cutoff_date and sched[d]["total"] > 0), default=None)
        ver = verify_loan(db, loan.id) or {}
        vals = [loan.loan_code, dorm.dormitory_name if dorm else "", dorm.school_name if dorm else "", fund.fund_name if fund else "",
                loan.assigned_amount, f"V{v.version_no}", sum(x["principal"] for x in sched.values()),
                sum(x["interest"] for x in sched.values()), remain, nxt, sched[nxt]["total"] if nxt else None,
                len(past), n_pay, short, "정상" if ver.get("overall") == "NORMAL" else "검토 필요", None, None]
        for ci, val in enumerate(vals, start=1):
            ws.cell(row=r, column=ci, value=val)
        ws.cell(row=r, column=17, value=f'=IF(P{r}="","",I{r}-P{r})')
        for col in (5, 7, 8, 9, 11, 16, 17):
            ws.cell(row=r, column=col).number_format = "#,##0"
        ws.cell(row=r, column=10).number_format = "yyyy-mm-dd"
        r += 1
    widths = [22, 16, 14, 16, 16, 10, 16, 16, 18, 14, 16, 12, 10, 12, 10, 20, 14]
    for ci, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    info = wb.create_sheet("이관 정보")
    for i, (k, val) in enumerate([("파일", batch.file_name), ("이관 기준일", str(batch.cutoff_date)),
                                  ("기준일 이전 납부완료 처리", "예" if batch.mark_paid else "아니오"),
                                  ("이관 일시", batch.created_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
                                  ("요약", ", ".join(f"{k}={v}" for k, v in (batch.summary or {}).items()))], start=1):
        info.cell(row=i, column=1, value=k).font = Font(bold=True)
        info.cell(row=i, column=2, value=val)
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 100
    info.cell(row=7, column=1, value="※ '기존 엑셀 잔여원금' 칸에 기존 관리 엑셀의 같은 기준일 잔여원금을 넣으면 '차이'가 자동 계산됩니다. 0이 아니면 조건·금리이력·중도변수를 확인하세요.")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cancel_batch(db: Session, batch: MigrationBatch, reason: str, user: User) -> dict:
    """이관 취소 = 이 이관으로 만든 차입금·기숙사를 숨김 삭제(관리자 › 삭제·복구에서 하나씩 되살릴 수 있음)."""
    from app.domain import deletion_service as dsvc
    reason = (reason or "").strip()
    if not reason:
        raise dsvc.DeletionError("취소 사유를 입력하세요.", 422)
    if batch.cancelled_at:
        raise dsvc.DeletionError("이미 취소한 이관입니다.")
    n_loans = n_dorms = 0
    alive = [l for l in (db.get(Loan, uuid.UUID(x)) for x in batch.loan_ids or []) if l is not None and not l.is_deleted]
    dsvc._ensure_no_pending(db, alive)          # 결재 대기 건이 있으면 하나도 지우지 않고 멈춤
    for lid in batch.loan_ids or []:
        loan = db.get(Loan, uuid.UUID(lid))
        if loan is not None and not loan.is_deleted:
            dsvc.delete_loan(db, loan.id, f"이관 취소({batch.file_name}): {reason}", user)
            n_loans += 1
    for did in batch.dormitory_ids or []:
        d = db.get(Dormitory, uuid.UUID(did))
        left = db.execute(select(Loan.id).where(Loan.dormitory_id == uuid.UUID(did), Loan.is_deleted.is_(False))).first()
        if d is not None and not d.is_deleted and not left:
            dsvc.delete_dormitory(db, d.id, f"이관 취소({batch.file_name}): {reason}", user)
            n_dorms += 1
    batch.cancelled_at = datetime.now(timezone.utc)
    batch.cancelled_by = user.id
    batch.cancel_reason = reason[:300]
    db.commit()
    return {"loans": n_loans, "dormitories": n_dorms}
