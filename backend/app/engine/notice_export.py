"""
안내공문 Excel/PDF 생성 — 첨부2 양식(6개 항목)을 그대로 재현한다.
공문은 '발급 당시 스냅샷'(GeneratedDocument.generated_data)에서 만든다. 그래서 같은 공문을 몇 번을
내려받아도 항상 같고, 이후 원리금 테이블이 바뀌어도 이미 발송한 공문 내용이 변하지 않는다.
원본과 같이 파생 칸(이자기간·상환후잔액·원리금합계·합계·검증)은 수식으로 둔다.
"""
import io
import shutil
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

FONT = "맑은 고딕"
MONEY = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'
DATE_FMT = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
IPMT_FILL = PatternFill("solid", fgColor="C6E0B4")      # 초록 음영(IPMT 적용 융자 건)
TOTAL_FILL = PatternFill("solid", fgColor="92D050")     # 합계행(원본과 동일색)
_THIN = Side(style="thin")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class PdfUnavailableError(Exception):
    pass


def _cell(ws, ref, value=None, *, bold=False, size=11, h="center", fmt=None, fill=None, border=True, wrap=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name=FONT, size=size, bold=bold)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    return c


def _merge(ws, rng, **kw):
    ws.merge_cells(rng)
    first = rng.split(":")[0]
    _cell(ws, first, **kw)
    # 병합 영역 테두리
    for row in ws[rng]:
        for c in row:
            c.border = BORDER


def build_notice_workbook(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    n = len(data["payments"])

    ws.column_dimensions["A"].width = 1.6
    ws.column_dimensions["B"].width = 17.4
    ws.column_dimensions["C"].width = 19.5
    ws.column_dimensions["D"].width = 16.5
    ws.column_dimensions["E"].width = 16.5
    ws.column_dimensions["F"].width = 16.5
    ws.column_dimensions["G"].width = 18.0
    ws.column_dimensions["H"].width = 2.0

    # 제목
    ws.row_dimensions[3].height = 26.25
    _merge(ws, "B3:G3", value=data["title"], bold=True, size=16)

    # 1. 납부기한
    _cell(ws, "B5", "1. 납부기한", bold=True, size=12, h="left", border=False)
    _cell(ws, "C5", date.fromisoformat(data["payment_deadline"]), fmt=DATE_FMT, border=False)

    # 2. 입금계좌
    _cell(ws, "B7", "2. 입금계좌", bold=True, size=12, h="left", border=False)
    _merge(ws, "B8:F8", value="계좌번호", fill=HEADER_FILL)
    _cell(ws, "G8", "예금주", fill=HEADER_FILL)
    _merge(ws, "B9:F9", value=data["account"]["bank_account"])
    _cell(ws, "G9", data["account"]["account_holder"])

    # 3. 이자기간
    int_start = 13
    int_end = 12 + n
    _cell(ws, "B11", "3. 이자기간", bold=True, size=12, h="left", border=False)
    for col, title in zip("BCDE", ["구 분", "시작일", "종료일", "이자기간"]):
        _cell(ws, f"{col}12", title, fill=HEADER_FILL)
    _merge(ws, "F12:G12", value="비고", fill=HEADER_FILL)

    bal_sec = int_end + 2
    bal_hdr = bal_sec + 1
    bal_start = bal_hdr + 1
    bal_end = bal_start + n - 1
    pay_sec = bal_end + 2
    pay_hdr = pay_sec + 1
    pay_start = pay_hdr + 1
    pay_end = pay_start + n - 1
    total_row = pay_end + 1
    note_sec = total_row + 2
    note_start = note_sec + 1
    note_end = note_start + 2

    for i, p in enumerate(data["interest_periods"]):
        r = int_start + i
        fill = IPMT_FILL if p["ipmt"] else None
        _cell(ws, f"B{r}", p["label"], fill=fill)
        _cell(ws, f"C{r}", date.fromisoformat(p["start"]), fmt=DATE_FMT, fill=fill)
        _cell(ws, f"D{r}", date.fromisoformat(p["end"]), fmt=DATE_FMT, fill=fill)
        _cell(ws, f"E{r}", f"=D{r}-C{r}", fmt="General", fill=fill)
    _merge(ws, f"F{int_start}:G{int_end}", value=data["interest_periods"][0]["note"], wrap=True)

    # 4. 대출잔액
    _cell(ws, f"B{bal_sec}", "4. 대출잔액", bold=True, size=12, h="left", border=False)
    for col, title in zip("BCDEFG", ["구 분", "융자액", "상환 전 대출잔액", "금번 상환원금", "상환 후 대출잔액", "-"]):
        _cell(ws, f"{col}{bal_hdr}", title, fill=HEADER_FILL)
    for i, b in enumerate(data["balances"]):
        r, pr = bal_start + i, pay_start + i
        fill = IPMT_FILL if b["ipmt"] else None
        _cell(ws, f"B{r}", b["label"], fill=fill)
        _cell(ws, f"C{r}", b["assigned_amount"], fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"D{r}", f"=C{pr}", fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"E{r}", f"=E{pr}", fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"F{r}", f"=D{r}-E{r}", fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"G{r}", f'=IF(C{pr}-(E{pr})=F{r},"TRUE","FALSE")', bold=True, fill=fill)

    # 5. 상환금액
    _cell(ws, f"B{pay_sec}", "5. 상환금액", bold=True, size=12, h="left", border=False)
    ws.row_dimensions[pay_hdr].height = 28.5
    for col, title in zip("BCDEFG", ["구 분", "상환잔액\n(직전 상환일 기준)", "적용이율", "원 금", "이 자", "원리금 합계"]):
        _cell(ws, f"{col}{pay_hdr}", title, fill=HEADER_FILL, wrap=True)
    for i, p in enumerate(data["payments"]):
        r = pay_start + i
        fill = IPMT_FILL if p["ipmt"] else None
        _cell(ws, f"B{r}", p["label"], fill=fill)
        _cell(ws, f"C{r}", p["balance_before"], fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"D{r}", round(p["rate"] / 100, 8), fmt="0.000%", fill=fill)
        _cell(ws, f"E{r}", p["principal"], fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"F{r}", p["interest"], fmt=MONEY, fill=fill, h="right")
        _cell(ws, f"G{r}", f"=E{r}+F{r}", fmt=MONEY, fill=fill, h="right")
    _cell(ws, f"B{total_row}", "합 계", bold=True, fill=TOTAL_FILL)
    _cell(ws, f"C{total_row}", f"=SUM(C{pay_start}:C{pay_end})", bold=True, fmt=MONEY, fill=TOTAL_FILL, h="right")
    _cell(ws, f"D{total_row}", "-", bold=True, fill=TOTAL_FILL)
    for col in "EFG":
        _cell(ws, f"{col}{total_row}", f"=SUM({col}{pay_start}:{col}{pay_end})", bold=True, fmt=MONEY, fill=TOTAL_FILL, h="right")

    # 6. 비고
    _cell(ws, f"B{note_sec}", "6. 비고", bold=True, size=12, h="left", border=False)
    _merge(ws, f"B{note_start}:G{note_end}", value="\n".join(data["notes"]) or None, h="left", wrap=True)

    ws.print_area = f"A1:H{note_end + 1}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.8, bottom=0.8)
    ws.print_options.horizontalCentered = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def convert_xlsx_to_pdf(xlsx_bytes: bytes) -> bytes:
    """LibreOffice(headless)로 PDF 변환. 서버에 LibreOffice가 없으면 PdfUnavailableError."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise PdfUnavailableError("서버에 PDF 변환기(LibreOffice)가 설치되어 있지 않습니다. Excel로 내려받으세요.")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "notice.xlsx"
        src.write_bytes(xlsx_bytes)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True, capture_output=True, timeout=120,
                env={"HOME": tmp, "PATH": "/usr/bin:/bin"},
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            raise PdfUnavailableError(f"PDF 변환에 실패했습니다: {e}")
        out = Path(tmp) / "notice.pdf"
        if not out.exists():
            raise PdfUnavailableError("PDF 변환 결과 파일이 없습니다.")
        return out.read_bytes()
