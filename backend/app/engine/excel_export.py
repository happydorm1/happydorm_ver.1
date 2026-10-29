"""
원리금 테이블 Excel 다운로드 생성.
19번 요구사항의 컬럼 구성을 그대로 따른다.
화면 조회와 동일한 계산결과를 사용한다(별도 재계산 없음 — 2-1 원칙).
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.loan import Loan
from app.models.schedule import RepaymentScheduleVersion

HEADER = [
    "연번", "차입금 구분", "상환예정일", "이자계산시작", "이자계산끝", "일수",
    "원금상환 회차", "상환 전 원금", "원금상환비율", "원금상환액",
    "이자", "미지급이자", "이자상환액", "원리금상환액", "상환 후 잔액",
    "적용금리", "비고",
]


def build_schedule_workbook(loan: Loan, version: RepaymentScheduleVersion, notes: dict | None = None) -> bytes:
    """notes: {상환일: 메모} — 사용자 메모(Phase 17)는 비고 뒤에 '| 메모: …'로 붙인다."""
    notes = notes or {}
    wb = Workbook()
    ws = wb.active
    ws.title = loan.loan_code[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F2A4D")
    for col, title in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(version.rows, start=2):
        principal = row.principal_total          # 상환유예 회수원금 포함(잔액 계산과 일치)
        ratio = round(principal / loan.assigned_amount, 5) if loan.assigned_amount else 0
        ws.append([
            row.seq_no,
            loan.loan_code,
            row.due_date,
            row.interest_calc_start,
            row.interest_calc_end,
            row.days,
            row.installment_no,
            row.balance_before,
            ratio,
            principal,
            row.interest_amount,
            row.unpaid_interest_amount,
            row.interest_total,
            row.principal_and_interest_total,
            row.balance_after,
            float(row.applied_rate),
            (row.remark or row.row_type.value) + (f" | 메모: {notes[row.due_date]}" if row.due_date in notes else ""),
        ])

    for col_idx in range(1, len(HEADER) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
