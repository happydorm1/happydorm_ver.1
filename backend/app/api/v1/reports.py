"""
§30 Excel 다운로드 8종 — 화면·계산과 같은 저장값을 그대로 내보낸다(재계산 없음). 권한 범위(SPC 담당자는 자기 SPC)·다운로드 권한을 따른다.

GET /api/v1/reports/excel?kind=
  schedules_all          1. 전체 원리금 테이블(차입금마다 시트, 첫 시트는 목록)
  schedules_dormitory    2. 기숙사별 원리금 테이블        (dormitory_id 필수)
  schedules_fund         3. 기금별 원리금 테이블          (fund_id 필수)
  period_repayments      4. 특정 기간 상환내역            (date_from, date_to 필수)
  actual_payments        5. 실제 상환이력(예정·실제·차액·미납) (as_of 기준, date_from/date_to 선택)
  midterm_events         6. 중도변수 이력(적용·취소·결재)
  rate_history           7. 금리변동 이력
  notice_data            8. 기숙사별 안내문 데이터        (due_date 필수)
"""
import io
import uuid
from datetime import date
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import ensure_can_download, get_current_user
from app.domain.calendar_service import load_holidays
from app.domain.loan_filters import LoanFilters, filtered_loans
from app.domain.payment_service import PaymentError, payment_summary
from app.domain.schedule_amounts import scheduled_by_due_date
from app.domain.version_utils import NoScheduleError, get_current_version, read_cache
from app.engine.business_day import effective_due_date
from app.domain.attachment_service import notes_by_due_date, remark_with_note
from app.engine.excel_export import HEADER as SCHEDULE_HEADER
from app.engine.year_attribution import year_attributed_interest
from app.models.event import EventRequest, LoanEvent
from app.models.loan import InterestRateHistory
from app.models.organization import User

router = APIRouter(prefix="/reports", tags=["Excel 다운로드"])

KINDS = {
    "schedules_all": "전체원리금테이블", "schedules_dormitory": "기숙사별원리금테이블", "schedules_fund": "기금별원리금테이블",
    "period_repayments": "기간별상환내역", "actual_payments": "실제상환이력", "midterm_events": "중도변수이력",
    "rate_history": "금리변동이력", "notice_data": "안내문데이터",
}
EVENT_LABEL = {"DEFERRAL": "상환유예", "EARLY_REPAYMENT": "조기상환", "ACCELERATION": "원금가속상환", "REFINANCING": "대환대출",
               "DELINQUENCY": "연체", "RATE_CHANGE": "금리변동"}
HEAD_FILL, HEAD_FONT = PatternFill("solid", fgColor="0F2A4D"), Font(bold=True, color="FFFFFF")
MONEY = "#,##0"


def _sheet(wb, title, header, rows, money_cols=(), widths=None):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for i in range(1, len(header) + 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(horizontal="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = MONEY
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
    ws.freeze_panes = "A2"
    return ws


def _schedule_rows(loan, version, notes=None):
    notes = notes or {}
    out = []
    ordered = sorted(version.rows, key=lambda x: x.seq_no)
    for r, att in zip(ordered, year_attributed_interest(ordered)):
        p = r.principal_total
        out.append([r.seq_no, loan.loan_code, r.due_date, r.interest_calc_start, r.interest_calc_end, r.days, r.installment_no,
                    r.balance_before, round(p / loan.assigned_amount, 5) if loan.assigned_amount else 0, p, r.interest_amount,
                    r.unpaid_interest_amount, r.interest_total, r.principal_and_interest_total, r.balance_after, float(r.applied_rate),
                    remark_with_note(r.remark or r.row_type.value, notes.get((loan.id, r.due_date))), att])
    return out


def build_report(db: Session, user: User, kind: str, *, dormitory_id=None, fund_id=None, date_from=None, date_to=None,
                 as_of: date, due_date=None) -> tuple[bytes, str]:
    f = LoanFilters(dormitory_id=dormitory_id, fund_id=fund_id)
    loans = filtered_loans(db, user, f)
    wb = Workbook()
    wb.remove(wb.active)
    holidays = load_holidays(db)
    with read_cache(db, [loan.id for loan, _, _ in loans]):
        if kind.startswith("schedules_"):
            notes = notes_by_due_date(db, [loan.id for loan, _, _ in loans])
            index = []
            for loan, dorm, fund in loans:
                try:
                    v = get_current_version(db, loan.id)
                except NoScheduleError:
                    continue
                index.append([dorm.dormitory_name, dorm.school_name, loan.loan_code, fund.fund_name, loan.assigned_amount,
                              f"V{v.version_no}", v.generated_reason, "정상" if v.status.value == "NORMAL" else "검토 필요"])
            _sheet(wb, "목록", ["기숙사", "학교", "차입금", "재원", "배정금액", "버전", "생성사유", "상태"], index, money_cols=(5,),
                   widths={1: 18, 2: 18, 3: 24})
            for loan, dorm, fund in loans:
                try:
                    v = get_current_version(db, loan.id)
                except NoScheduleError:
                    continue
                name = f"{dorm.dormitory_name[:12]}_{loan.loan_code[-12:]}"
                _sheet(wb, name, SCHEDULE_HEADER, _schedule_rows(loan, v, notes), money_cols=(8, 10, 11, 12, 13, 14, 15, 18))
        elif kind == "period_repayments":
            rows = []
            for loan, dorm, fund in loans:
                try:
                    v = get_current_version(db, loan.id)
                except NoScheduleError:
                    continue
                for d, m in sorted(scheduled_by_due_date(db, v).items()):
                    if date_from <= d <= date_to and m["total"] > 0:
                        rows.append([dorm.dormitory_name, loan.loan_code, fund.fund_name, d, effective_due_date(d, holidays),
                                     m["principal"], m["interest"], m["total"], f"V{v.version_no}"])
            rows.sort(key=lambda r: (r[3], r[0], r[1]))
            rows.append(["합계", "", "", "", "", sum(r[5] for r in rows), sum(r[6] for r in rows), sum(r[7] for r in rows), ""])
            _sheet(wb, f"{date_from}~{date_to}", ["기숙사", "차입금", "재원", "상환예정일", "납부기한", "원금", "이자", "원리금", "버전"],
                   rows, money_cols=(6, 7, 8), widths={1: 18, 2: 24})
        elif kind == "actual_payments":
            label = {"PAID": "완납", "OVERPAID": "초과납부", "PARTIAL": "일부납부", "UNPAID": "미납", "PENDING": "납부기한 전"}
            rows = []
            for loan, dorm, fund in loans:
                try:
                    s = payment_summary(db, loan, as_of, holidays)
                except PaymentError:
                    continue
                for r in s["rows"]:
                    if (date_from and r["due_date"] < date_from) or (date_to and r["due_date"] > date_to):
                        continue
                    if r["status"] == "PENDING" and not r["paid_total"]:
                        continue
                    rows.append([dorm.dormitory_name, loan.loan_code, fund.fund_name, r["due_date"], r["payment_deadline"],
                                 r["scheduled_principal"], r["scheduled_interest"], r["scheduled_total"], r["last_paid_date"],
                                 r["paid_principal"], r["paid_interest"], r["paid_total"], r["difference"], r["outstanding"],
                                 label.get(r["status"], r["status"])])
            _sheet(wb, f"기준일 {as_of}", ["기숙사", "차입금", "재원", "상환예정일", "납부기한", "예정원금", "예정이자", "예정원리금",
                                         "실제 납입일", "실제원금", "실제이자", "실제원리금", "차액", "미납금액", "상태"],
                   rows, money_cols=(6, 7, 8, 10, 11, 12, 13, 14), widths={1: 18, 2: 24})
        elif kind == "midterm_events":
            ids = {loan.id: (loan, dorm) for loan, dorm, _ in loans}
            names = {u.id: u.name for u in db.execute(select(User)).scalars()}
            rows = []
            for e in db.execute(select(LoanEvent).where(LoanEvent.loan_id.in_(ids)).order_by(LoanEvent.event_date)).scalars() if ids else []:
                loan, dorm = ids[e.loan_id]
                rows.append([dorm.dormitory_name, loan.loan_code, EVENT_LABEL.get(e.event_type.value, e.event_type.value), e.event_date,
                             "적용" if e.status.value == "APPLIED" else "취소", e.reason, names.get(e.created_by), e.created_at.replace(tzinfo=None)])
            for r in db.execute(select(EventRequest).where(EventRequest.loan_id.in_(ids)).order_by(EventRequest.created_at)).scalars() if ids else []:
                loan, dorm = ids[r.loan_id]
                rows.append([dorm.dormitory_name, loan.loan_code, EVENT_LABEL.get(r.event_type, r.event_type) + "(기안)",
                             r.payload.get("event_date") or r.payload.get("effective_start_date"),
                             {"REQUESTED": "결재 대기", "APPROVED": "승인", "REJECTED": "반려", "WITHDRAWN": "회수"}[r.status.value],
                             r.decision_note or r.payload.get("reason"), names.get(r.requested_by), r.created_at.replace(tzinfo=None)])
            _sheet(wb, "중도변수", ["기숙사", "차입금", "구분", "적용일", "상태", "사유·결재의견", "등록자", "등록일시"], rows,
                   widths={1: 18, 2: 24, 6: 30, 8: 20})
        elif kind == "rate_history":
            ids = {loan.id: (loan, dorm, fund) for loan, dorm, fund in loans}
            rows = []
            for h in db.execute(select(InterestRateHistory).where(InterestRateHistory.loan_id.in_(ids))
                                .order_by(InterestRateHistory.loan_id, InterestRateHistory.effective_start_date)).scalars() if ids else []:
                loan, dorm, fund = ids[h.loan_id]
                rows.append([dorm.dormitory_name, loan.loan_code, fund.fund_name, h.effective_start_date, h.effective_end_date,
                             float(h.annual_rate), h.rate_table_name, h.rate_basis, "사용중지" if h.is_deleted else "사용",
                             h.created_at.replace(tzinfo=None)])
            _sheet(wb, "금리이력", ["기숙사", "차입금", "재원", "적용시작일", "적용종료일", "금리(%)", "금리적용표", "산출근거", "상태", "등록일시"],
                   rows, widths={1: 18, 2: 24, 10: 20})
        elif kind == "notice_data":
            rows = []
            for loan, dorm, fund in loans:
                try:
                    v = get_current_version(db, loan.id)
                except NoScheduleError:
                    continue
                m = scheduled_by_due_date(db, v).get(due_date)
                if not m or m["total"] <= 0:
                    continue
                day_rows = sorted((r for r in v.rows if r.due_date == due_date), key=lambda r: r.seq_no)
                start, end = min(r.interest_calc_start for r in day_rows), max(r.interest_calc_end for r in day_rows)
                rows.append([dorm.dormitory_name, dorm.school_name, loan.loan_code, fund.fund_name, start, end, (end - start).days,
                             loan.assigned_amount, day_rows[0].balance_before, m["principal"], day_rows[-1].balance_after,
                             float(day_rows[-1].applied_rate), m["interest"], m["total"], effective_due_date(due_date, holidays)])
            _sheet(wb, f"{due_date} 상환분", ["기숙사", "학교", "차입금", "재원", "이자시작(초일산입)", "이자끝(말일불산입)", "일수", "배정금액",
                                             "상환전잔액", "원금", "상환후잔액", "이율(%)", "이자", "합계", "납부기한"],
                   rows, money_cols=(8, 9, 10, 11, 13, 14), widths={1: 18, 2: 18, 3: 24})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{KINDS[kind]}_{as_of}.xlsx"


@router.get("/excel")
def download_report(kind: str, dormitory_id: uuid.UUID | None = None, fund_id: uuid.UUID | None = None,
                    date_from: date | None = None, date_to: date | None = None, due_date: date | None = None,
                    as_of: date | None = None, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    if kind not in KINDS:
        raise HTTPException(status_code=422, detail="kind는 " + ", ".join(KINDS) + " 중 하나입니다.")
    need = {"schedules_dormitory": ("dormitory_id", dormitory_id), "schedules_fund": ("fund_id", fund_id),
            "notice_data": ("due_date", due_date)}.get(kind)
    if need and need[1] is None:
        raise HTTPException(status_code=422, detail=f"{need[0]}가 필요합니다.")
    if kind == "period_repayments" and not (date_from and date_to):
        raise HTTPException(status_code=422, detail="date_from과 date_to가 필요합니다.")
    if date_from and date_to and date_from > date_to:
        raise HTTPException(status_code=422, detail="기간 시작일이 종료일보다 늦습니다.")
    ensure_can_download(current_user)
    content, name = build_report(db, current_user, kind, dormitory_id=dormitory_id, fund_id=fund_id, date_from=date_from,
                                 date_to=date_to, as_of=as_of or date.today(), due_date=due_date)
    return Response(content=content, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})
