"""
Phase 15 최종 검증 — 원본 엑셀(첨부)과 프로그램 결과를 행 단위로 다시 대조해 결과 엑셀을 만든다(§38, §51 1~3순위).

    DATABASE_URL=...(테스트용 DB) python -m verification.excel_reconciliation --out Excel대조결과.xlsx

실제 API + DB로 차입금을 만들고 계산한 뒤 첨부 파일의 값과 칸마다 비교한다. 운영 DB에는 실행하지 않는다(검증용 차입금이 생긴다).
첨부 파일 위치: --uploads (기본 /mnt/user-data/uploads)
"""
import argparse
import sys
import uuid
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

UP = Path("/mnt/user-data/uploads")


def _files(up: Path):
    return {k: up / v for k, v in {
        "x1": "첨부_1__원리금_테이블.xlsx", "x2": "첨부_2__기숙사별_안내_공문.xlsx", "x4": "첨부_4__원금가속상환_사례_테이블.xlsx",
        "x6": "첨부_6__일부대환_사례_테이블.xlsx", "x7": "첨부_7__전액대환_사례_테이블.xlsx"}.items()}


class Case:
    def __init__(self, key, title, source, criteria):
        self.key, self.title, self.source, self.criteria = key, title, source, criteria
        self.rows: list[dict] = []          # 비교 행
        self.known: set = set()             # 의도적 예외(수동보정 미재현) 행 번호
        self.note = ""

    def add(self, no, label, item, program, excel, exception=False):
        ok = program == excel
        self.rows.append({"no": no, "label": label, "item": item, "program": program, "excel": excel,
                          "result": "일치" if ok else ("예외(원본 수동보정)" if exception else "불일치")})

    @property
    def counts(self):
        c = {"일치": 0, "불일치": 0, "예외(원본 수동보정)": 0}
        for r in self.rows:
            c[r["result"]] += 1
        return c


def _d(v):
    return v.date().isoformat() if isinstance(v, datetime) else (v.isoformat() if isinstance(v, date) else v)


# ---------------- A. 공공자금관리기금(원금균등) ----------------
def case_a(F):
    import test_phase3_schedule as t3
    c = Case("A", "테스트 A — 2015-08-21 공공자금관리기금(원금균등, 분기)", "첨부1 '2015-08-21_공자' 13~173행",
             "회차마다 이자일수·원금·이자(연도경계 분할행 포함)")
    h = t3._admin_headers()
    run = uuid.uuid4().hex[:8]
    dorm = t3._setup_dormitory(h, run)
    fund, q, ep = t3._fund(h, "공공자금관리기금"), t3._code(h, "REPAY_CYCLE", "QUARTERLY"), t3._code(h, "REPAY_METHOD", "EQUAL_PRINCIPAL")
    r = t3.client.post("/api/v1/loans", json={
        "dormitory_id": dorm, "fund_id": fund["id"], "loan_code": f"VERIFY-A-{run}", "assigned_amount": 1_768_000_000,
        "assigned_date": "2015-08-21", "rate_type": "VARIABLE", "repay_cycle_code_id": q["id"], "repay_method_code_id": ep["id"],
        "grace_start_date": "2015-08-21", "grace_end_date": "2022-08-31", "principal_repay_start_date": "2022-11-30",
        "repay_period_months": 300, "initial_rate": {"effective_start_date": "2015-08-21", "effective_end_date": "2022-02-28",
                                                     "annual_rate": "2.24600"}}, headers=h)
    loan = r.json()["id"]
    t3.client.post(f"/api/v1/loans/{loan}/rate-history", json={"effective_start_date": "2022-02-28", "annual_rate": "2.32000"}, headers=h)
    g = t3.client.post(f"/api/v1/loans/{loan}/schedule/generate", headers=h).json()["rows"]
    ws = openpyxl.load_workbook(F["x1"], data_only=True)["2015-08-21_공자"]
    for i, row in enumerate(range(13, 174)):
        x = g[i]
        second = i > 0 and g[i - 1]["due_date"] == x["due_date"]
        k, l_ = ws.cell(row=row, column=11).value or 0, ws.cell(row=row, column=12).value or 0
        exp_int = k if (x["row_type"] != "YEAR_END_SPLIT_REAL" or second) else l_
        lab = f"{x['due_date']} ({x['installment_no'] or '분할/거치'})"
        c.add(i + 1, lab, "이자일수", x["days"], ws.cell(row=row, column=7).value)
        c.add(i + 1, lab, "원금", x["principal_amount"], ws.cell(row=row, column=10).value or 0)
        c.add(i + 1, lab, "이자", x["interest_amount"], exp_int)
    c.add(len(g), "합계", "총 원금상환액", sum(x["principal_amount"] for x in g), 1_768_000_000)
    c.add(len(g), "최종", "최종 잔여원금", g[-1]["balance_after"], 0)
    return c


# ---------------- B. 주택도시기금(원리금균등) ----------------
def case_b(F):
    import test_phase3_schedule as t3
    c = Case("B", "테스트 B — 2016-03-28 주택도시기금(원리금균등, 반기)", "첨부1 '2016-03-28_주택' 13~102행",
             "회차마다 원금·이자. 52·54회차는 원본의 수동보정(-3원/+4원)을 재현하지 않고 '누적원금·최종잔액 기준 자동보정'으로 처리(§39)")
    h = t3._admin_headers()
    run = uuid.uuid4().hex[:8]
    dorm = t3._setup_dormitory(h, run)
    fund, s, ep = t3._fund(h, "주택도시기금"), t3._code(h, "REPAY_CYCLE", "SEMIANNUAL"), t3._code(h, "REPAY_METHOD", "EQUAL_PAYMENT")
    loan = t3.client.post("/api/v1/loans", json={
        "dormitory_id": dorm, "fund_id": fund["id"], "loan_code": f"VERIFY-B-{run}", "assigned_amount": 4_097_000_000,
        "assigned_date": "2016-03-28", "rate_type": "FIXED", "repay_cycle_code_id": s["id"], "repay_method_code_id": ep["id"],
        "grace_start_date": "2016-03-28", "grace_end_date": "2019-02-28", "principal_repay_start_date": "2019-08-31",
        "repay_period_months": 324, "initial_rate": {"effective_start_date": "2016-03-28", "annual_rate": "2.00000"}}, headers=h).json()["id"]
    g = t3.client.post(f"/api/v1/loans/{loan}/schedule/generate", headers=h).json()["rows"]
    ws = openpyxl.load_workbook(F["x1"], data_only=True)["2016-03-28_주택"]
    for i, row in enumerate(range(13, 103)):
        x = g[i]
        exc = x["installment_no"] in (52, 54)
        lab = f"{x['due_date']} ({x['installment_no'] or '거치'})"
        c.add(i + 1, lab, "원금", x["principal_amount"], ws.cell(row=row, column=10).value or 0, exception=exc)
        c.add(i + 1, lab, "이자", x["interest_amount"], ws.cell(row=row, column=13).value or 0, exception=exc)
    c.add(len(g), "합계", "총 원금상환액", sum(x["principal_amount"] for x in g), 4_097_000_000)
    c.add(len(g), "최종", "최종 잔여원금", g[-1]["balance_after"], 0)
    c.note = "원본 52회차 원금은 PPMT-3원, 54회차는 +4원으로 손으로 고친 값. 프로그램은 같은 합계를 마지막 회차 자동보정으로 맞춘다."
    return c


# ---------------- C. 상환유예 ----------------
def case_c(F):
    import test_phase6_deferral as t6
    c = Case("C", "상환유예 — 2014-12-18 주택도시기금, 6회차 원금+이자 동결, 분기 10회 회수",
             "첨부3 값(원본 파일은 이번 세션에 없음 — 2단계에서 첨부3으로 확정해 테스트에 고정한 값)",
             "동결 원금·이자, 회수 시작일, 회수 10회의 원금·이자·유예원금 추가이자")
    h = t6._admin_headers()
    run = uuid.uuid4().hex[:8]
    cl = t6.client
    spc = cl.post("/api/v1/spc", json={"spc_name": f"V-{run}"}, headers=h).json()["id"]
    dorm = cl.post("/api/v1/dormitories", json={"dormitory_name": f"검증-{run}", "school_name": "검증대", "spc_id": spc}, headers=h).json()["id"]
    funds = {f["fund_name"]: f for f in cl.get("/api/v1/reference/funds", headers=h).json()}
    cyc = {x["code"]: x for x in cl.get("/api/v1/reference/codes", params={"group_code": "REPAY_CYCLE"}, headers=h).json()}
    mth = {x["code"]: x for x in cl.get("/api/v1/reference/codes", params={"group_code": "REPAY_METHOD"}, headers=h).json()}
    loan = cl.post("/api/v1/loans", json={
        "dormitory_id": dorm, "fund_id": funds["주택도시기금"]["id"], "loan_code": f"VERIFY-C-{run}", "assigned_amount": 5_517_000_000,
        "assigned_date": "2014-12-18", "rate_type": "FIXED", "repay_cycle_code_id": cyc["SEMIANNUAL"]["id"],
        "repay_method_code_id": mth["EQUAL_PAYMENT"]["id"], "grace_start_date": "2014-12-18", "grace_end_date": "2017-08-31",
        "principal_repay_start_date": "2018-02-28", "repay_period_months": 324,
        "initial_rate": {"effective_start_date": "2014-12-18", "annual_rate": "2.00000"}}, headers=h).json()["id"]
    cl.post(f"/api/v1/loans/{loan}/schedule/generate", headers=h)
    d = cl.post(f"/api/v1/loans/{loan}/events/deferral", json={"event_date": "2020-08-31", "defer_target": "PRINCIPAL_INTEREST",
                                                                "recovery_cycle_months": 3, "recovery_count": 10}, headers=h).json()["event"]["deferral"]
    c.add(0, "동결", "동결 원금", d["frozen_principal_total"], 81_506_004)
    c.add(0, "동결", "동결 이자", d["frozen_interest_total"], 51_214_162)
    c.add(0, "회수", "회수 시작일", d["recovery_start_date"], "2021-08-31")
    ep = [8_150_600] * 9 + [8_150_604]
    ei = [5_121_416] * 9 + [5_121_418]
    ea = [82176, 40641, 40195, 41088, 41088, 40641, 40195, 41088, 41088, 40641]
    for i, r in enumerate(d["recovery_rows"]):
        lab = f"회수 {i + 1}회 {r.get('due_date', '')}"
        c.add(i + 1, lab, "회수 원금", r["recovery_principal"], ep[i])
        c.add(i + 1, lab, "회수 이자", r["recovery_interest"], ei[i])
        c.add(i + 1, lab, "유예원금 추가이자", r["additional_interest_on_frozen_principal"], ea[i])
    c.note = "첨부3 원본 파일을 다시 올려 주시면 원본 칸과 직접 대조할 수 있습니다."
    return c


# ---------------- D. 원금가속상환 ----------------
def case_d(F):
    import test_phase8_events as t8
    c = Case("D", "원금가속상환 — 2019-12-30 주택도시기금, 2026-11-30부터 가속배율 1.5", "첨부4 '2019-12-30_주택(가속상환)'",
             "가속 이후 회차(34~71행)의 상환전잔액·원금·이자")
    h = t8._headers()
    loan = t8._make_loan(h, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT", start="2019-12-30", first="2023-02-28",
                         months=324, amount=15_055_000_000, rate="1.57200", grace=("2019-12-30", "2022-08-31"))
    v2 = t8.client.post(f"/api/v1/loans/{loan}/events/acceleration", json={"event_date": "2026-11-30", "acceleration_multiplier": "1.5"},
                        headers=h).json()
    ws = openpyxl.load_workbook(F["x4"], data_only=True)["2019-12-30_주택(가속상환)"]
    for i, x in enumerate(v2["rows"]):
        rr = 13 + i
        if i < 33:
            continue
        lab = f"{x['due_date']} ({x['installment_no'] or '분할'})"
        exp_int = (ws.cell(row=rr, column=13).value or 0) if x["installment_no"] is None else (ws.cell(row=rr, column=12).value or 0)
        c.add(i + 1, lab, "상환전잔액", x["balance_before"], ws.cell(row=rr, column=9).value)
        c.add(i + 1, lab, "원금", x["principal_amount"], ws.cell(row=rr, column=11).value or 0)
        c.add(i + 1, lab, "이자", x["interest_amount"], exp_int)
    c.add(len(v2["rows"]), "최종", "최종 잔여원금", v2["rows"][-1]["balance_after"], 0)
    return c


# ---------------- E/F. 대환 ----------------
def _refi_case(F, partial: bool):
    import test_phase8_refinancing as tr
    h = tr._headers()
    if partial:
        c = Case("E", "일부대환 — 2010-12-22 공공자금관리기금 → 주택도시기금 35억 전환, 2027-09-30 재합류", "첨부6 '2010-12-22_공자(일부 대환 예정)'",
                 "대환 구간 이후 모든 행의 상환일·이자일수·잔액·원금·이자")
        S = tr._read_sheet(F["x6"], "2010-12-22_공자(일부 대환 예정)")
        win = next(i for i, x in enumerate(S) if x["fund"] == "주택도시기금")
        rej = next(i for i, x in enumerate(S) if i > win and x["fund"] == "공공자금관리기금" and S[i - 1]["fund"] == "공공자금관리기금"
                   and x["due"] >= date(2027, 9, 30))
        before, window, after = S[:win], S[win:rej], S[rej:]
        conv = [x for x in window if x["fund"] == "주택도시기금"]
        single = list(before) + [dict(x, fund="합산") for x in conv] + list(after)
        ledger, bal, inst = [], 13_740_000_000, 0
        for i, x in enumerate(single):
            memo = i + 1 < len(single) and single[i + 1]["due"] == x["due"]
            bb = x["bal"] if (x in before or x in after) else bal
            pr = 0 if memo else x["J"]
            if not memo:
                inst += 1
            ledger.append((x["due"], x["start"], x["end"], x["days"], None if memo else inst, bb, pr, x["K"] or x["L"], memo))
            bal = bb - pr
        loan, funds = tr._loan_with_ledger(h, ledger, assigned=13_740_000_000, assigned_date="2010-12-22", rate=4.557)
        v2 = tr.client.post(f"/api/v1/loans/{loan}/events/refinancing", json={
            "refi_type": "PARTIAL", "event_date": "2025-03-31", "new_fund_id": funds["주택도시기금"]["id"], "new_annual_rate": "2.17000",
            "converted_amount": 3_500_000_000, "remaining_amount": 4_658_000_000, "rejoin_date": "2027-09-30"}, headers=h).json()
        start = win
    else:
        c = Case("F", "전액대환 — 2009-03-09 주택도시기금 같은 원장 승계, 2024-06-07 신금리 2.345%", "첨부7 '2009-03-09_주택(대환)'",
                 "대환 이후 모든 행의 상환일·이자일수·잔액·원금·이자")
        S = tr._read_sheet(F["x7"], "2009-03-09_주택(대환)", last=300)
        ev = next(i for i, x in enumerate(S) if x["row"] == 88)
        ledger, inst = [], 0
        for i, x in enumerate(S):
            memo = i + 1 < len(S) and S[i + 1]["due"] == x["due"]
            if not memo:
                inst += 1
            ledger.append((x["due"], x["start"], x["end"], x["days"], None if memo else inst, x["bal"], 0 if memo else x["J"],
                           x["K"] or x["L"], memo))
        loan, funds = tr._loan_with_ledger(h, ledger, assigned=7_500_000_000, assigned_date="2009-03-09", rate=4.643)
        v2 = tr.client.post(f"/api/v1/loans/{loan}/events/refinancing", json={
            "refi_type": "FULL", "event_date": "2024-06-07", "new_fund_id": funds["주택도시기금"]["id"], "new_annual_rate": "2.34500"},
            headers=h).json()
        start = ev + 1
    for gi, (g, a) in enumerate(zip(v2["rows"], S)):
        if gi < start:
            continue
        lab = f"{g['due_date']} (원본 {a['row']}행)"
        c.add(gi + 1, lab, "상환일", g["due_date"], a["due"].isoformat())
        c.add(gi + 1, lab, "이자일수", g["days"], a["days"])
        c.add(gi + 1, lab, "상환전잔액", g["balance_before"], a["bal"])
        c.add(gi + 1, lab, "원금", g["principal_amount"], a["J"])
        c.add(gi + 1, lab, "이자", g["interest_amount"], a["K"] if a["K"] else a["L"])
    c.add(len(v2["rows"]), "행 수", "전체 행 수", len(v2["rows"]), len(S))
    c.add(len(v2["rows"]), "최종", "최종 잔여원금", v2["rows"][-1]["balance_after"], 0)
    return c


# ---------------- G. 안내공문 ----------------
def case_g(F):
    import test_phase11_notice as tn
    from openpyxl import load_workbook
    c = Case("G", "안내공문 — 2026-08-31 상환분(공자 원금균등 + 주택 원리금균등 2건)", "첨부2 Sheet1 (이자기간·대출잔액·상환금액)",
             "이자기간, 배정·상환전·원금·상환후 잔액, 이율, 원금, 이자")
    h = tn.admin_headers()
    dorm, a, b = tn._nazarene_dormitory(h)
    tn._template(h)
    d = tn.client.get(f"/api/v1/dormitories/{dorm}/notices/preview", params={"due_date": "2026-08-31"}, headers=h).json()
    ws2 = load_workbook(F["x2"], data_only=True)["Sheet1"]
    for i, row in enumerate((13, 14)):
        p = d["interest_periods"][i]
        c.add(i + 1, f"이자기간 {p['label']}", "시작일", p["start"], _d(ws2[f"C{row}"].value))
        c.add(i + 1, f"이자기간 {p['label']}", "종료일", p["end"], _d(ws2[f"D{row}"].value))
    for i, row in enumerate((18, 19)):
        bl = d["balances"][i]
        lab = f"대출잔액 {('공공자금관리기금', '주택도시기금')[i]}"
        for item, key, col in (("배정금액", "assigned_amount", "C"), ("상환전잔액", "balance_before", "D"),
                               ("원금", "principal", "E"), ("상환후잔액", "balance_after", "F")):
            c.add(i + 1, lab, item, bl[key], ws2[f"{col}{row}"].value)
    for i, row in enumerate((23, 24)):
        p = d["payments"][i]
        lab = f"상환금액 {('공공자금관리기금', '주택도시기금')[i]}"
        c.add(i + 1, lab, "이율(%)", round(p["rate"], 6), round(ws2[f"D{row}"].value * 100, 6))
        c.add(i + 1, lab, "원금", p["principal"], ws2[f"E{row}"].value)
        c.add(i + 1, lab, "이자", p["interest"], ws2[f"F{row}"].value, exception=(i == 1))
    c.note = ("주택도시기금 이자: 첨부2 공문은 32,361,800원, 원리금 테이블(첨부1) 원본 값은 32,361,807원. "
              "공문을 손으로 옮겨 적으며 생긴 7원 차이로, 프로그램은 원리금 테이블 값을 그대로 연결한다.")
    return c


# ---------------- 결과 엑셀 ----------------
HEAD = PatternFill("solid", fgColor="12325A")
BAD = PatternFill("solid", fgColor="F8D7D7")
EXC = PatternFill("solid", fgColor="FFF1CC")
OKF = PatternFill("solid", fgColor="E3F1E3")
THIN = Border(*(Side(style="thin", color="D0D0D0"),) * 4)


def _style_header(ws, row, n):
    for col in range(1, n + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD, Font(color="FFFFFF", bold=True), Alignment(horizontal="center"), THIN


def write(cases, out: Path):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "요약"
    s["A1"] = "원본 엑셀 대조 결과 (Phase 15 최종 검증)"
    s["A1"].font = Font(size=14, bold=True)
    s["A2"] = f"실행 {datetime.now():%Y-%m-%d %H:%M} · 실제 API+DB로 계산한 값과 첨부 엑셀 값을 칸마다 비교"
    hdr = ["구분", "대상", "원본", "비교 항목", "비교 칸 수", "일치", "예외(원본 수동보정)", "불일치", "판정", "비고"]
    s.append([])
    s.append(hdr)
    _style_header(s, 4, len(hdr))
    for c in cases:
        n = c.counts
        verdict = "통과" if n["불일치"] == 0 else "확인 필요"
        s.append([c.key, c.title, c.source, c.criteria, len(c.rows), n["일치"], n["예외(원본 수동보정)"], n["불일치"], verdict, c.note])
        r = s.max_row
        s.cell(row=r, column=9).fill = OKF if verdict == "통과" else BAD
    for col, w in zip("ABCDEFGHIJ", (6, 46, 36, 44, 10, 8, 12, 8, 9, 60)):
        s.column_dimensions[col].width = w
    for row in s.iter_rows(min_row=5, max_row=s.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN
    for c in cases:
        ws = wb.create_sheet(f"{c.key}_상세")
        ws.append([c.title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"원본: {c.source}"])
        ws.append([])
        ws.append(["행", "회차/구분", "항목", "프로그램", "원본 엑셀", "차이", "결과"])
        _style_header(ws, 4, 7)
        for r in c.rows:
            diff = (r["program"] - r["excel"]) if isinstance(r["program"], (int, float)) and isinstance(r["excel"], (int, float)) else ""
            ws.append([r["no"], r["label"], r["item"], r["program"], r["excel"], diff, r["result"]])
            cell = ws.cell(row=ws.max_row, column=7)
            cell.fill = OKF if r["result"] == "일치" else (EXC if r["result"].startswith("예외") else BAD)
            for col in (4, 5, 6):
                if isinstance(ws.cell(row=ws.max_row, column=col).value, int):
                    ws.cell(row=ws.max_row, column=col).number_format = "#,##0"
        for col, w in zip("ABCDEFG", (6, 30, 18, 18, 18, 10, 18)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A5"
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="Excel대조결과.xlsx")
    ap.add_argument("--uploads", default=str(UP))
    a = ap.parse_args()
    F = _files(Path(a.uploads))
    missing = [str(p) for p in F.values() if not p.exists()]
    if missing:
        raise SystemExit("첨부 파일이 없습니다: " + ", ".join(missing))
    cases = [case_a(F), case_b(F), case_c(F), case_d(F), _refi_case(F, True), _refi_case(F, False), case_g(F)]
    write(cases, Path(a.out))
    for c in cases:
        n = c.counts
        print(f"{c.key} {c.title[:40]:40} 칸 {len(c.rows):5} 일치 {n['일치']:5} 예외 {n['예외(원본 수동보정)']:2} 불일치 {n['불일치']}")
    print("저장:", a.out)


if __name__ == "__main__":
    main()
