"""Phase 11: 안내공문 자동생성 — 첨부2(실제 공문)와 첨부1(원리금 테이블) 대조"""
import io
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

import pytest
from _helpers import admin_headers, client, make_loan, make_org, refs, spc_manager_headers
from openpyxl import load_workbook

EXCEL1 = "/mnt/user-data/uploads/첨부_1__원리금_테이블.xlsx"
EXCEL2 = "/mnt/user-data/uploads/첨부_2__기숙사별_안내_공문.xlsx"
ACCOUNT = {"bank_account": "테스트은행 000-00-000000", "account_holder": "테스트기금"}   # 실제 계좌번호는 코드/테스트에 넣지 않는다


def _template(h):
    r = client.post("/api/v1/document-templates", json={"template_name": f"양식-{uuid.uuid4().hex[:6]}", "config": ACCOUNT}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _raw_loan(h, dorm, *, fund, cycle, method, amount, start, grace, first, months, rates):
    funds, cyc, mth = refs(h)
    r = client.post("/api/v1/loans", json={
        "dormitory_id": dorm, "fund_id": funds[fund]["id"], "loan_code": f"N-{uuid.uuid4().hex[:8]}",
        "assigned_amount": amount, "assigned_date": start, "rate_type": "VARIABLE",
        "repay_cycle_code_id": cyc[cycle]["id"], "repay_method_code_id": mth[method]["id"],
        "grace_start_date": grace[0], "grace_end_date": grace[1],
        "principal_repay_start_date": first, "repay_period_months": months,
        "initial_rate": {"effective_start_date": start, "effective_end_date": rates[0][1], "annual_rate": rates[0][2]}}, headers=h)
    assert r.status_code == 201, r.text
    loan_id = r.json()["id"]
    for s, e, rate in rates[1:]:
        assert client.post(f"/api/v1/loans/{loan_id}/rate-history",
                           json={"effective_start_date": s, "effective_end_date": e, "annual_rate": rate}, headers=h).status_code == 201
    assert client.post(f"/api/v1/loans/{loan_id}/schedule/generate", headers=h).status_code == 201
    return loan_id


def _nazarene_dormitory(h):
    """첨부2의 사례: 한 기숙사의 공공자금관리기금(원금균등) + 주택도시기금(원리금균등) 두 차입금."""
    _, dorm = make_org(h)
    a = _raw_loan(h, dorm, fund="공공자금관리기금", cycle="QUARTERLY", method="EQUAL_PRINCIPAL", amount=1_768_000_000,
                  start="2015-08-21", grace=("2015-08-21", "2022-08-31"), first="2022-11-30", months=300,
                  rates=[("2015-08-21", "2022-02-28", "2.24600"), ("2022-02-28", None, "2.32000")])
    b = _raw_loan(h, dorm, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT", amount=4_097_000_000,
                  start="2016-03-28", grace=("2016-03-28", "2019-02-28"), first="2019-08-31", months=324,
                  rates=[("2016-03-28", None, "2.00000")])
    return dorm, a, b


@pytest.mark.skipif(not (Path(EXCEL1).exists() and Path(EXCEL2).exists()), reason="첨부1·2 파일 필요")
def test_notice_matches_real_attachment_2():
    h = admin_headers()
    dorm, loan_a, loan_b = _nazarene_dormitory(h)
    _template(h)

    r = client.get(f"/api/v1/dormitories/{dorm}/notices/preview", params={"due_date": "2026-08-31"}, headers=h)
    assert r.status_code == 200, r.text
    d = r.json()

    ws2 = load_workbook(EXCEL2, data_only=True)["Sheet1"]
    def d2(ref):
        v = ws2[ref].value
        return v.date().isoformat() if hasattr(v, "date") else v

    # --- 3. 이자기간
    for i, row in enumerate((13, 14)):
        p = d["interest_periods"][i]
        assert (p["label"], p["start"], p["end"]) == (ws2[f"B{row}"].value, d2(f"C{row}"), d2(f"D{row}"))
        assert p["days"] == (ws2[f"D{row}"].value - ws2[f"C{row}"].value).days
    assert d["interest_periods"][0]["note"] == "초일 산입, 말일 불산입"

    # --- 4. 대출잔액 / 5. 상환금액 (공공자금관리기금: 첨부2와 전 항목 일치)
    b, p = d["balances"][0], d["payments"][0]
    assert b["assigned_amount"] == ws2["C18"].value
    assert (b["balance_before"], b["principal"], b["balance_after"]) == (ws2["D18"].value, ws2["E18"].value, ws2["F18"].value)
    assert p["rate"] == pytest.approx(ws2["D23"].value * 100)
    assert (p["principal"], p["interest"]) == (ws2["E23"].value, ws2["F23"].value)

    # --- 주택도시기금: 원금·잔액·이율 일치, 이자는 첨부1 원리금 테이블 값과 일치
    b, p = d["balances"][1], d["payments"][1]
    assert b["assigned_amount"] == ws2["C19"].value
    assert (b["balance_before"], b["principal"], b["balance_after"]) == (ws2["D19"].value, ws2["E19"].value, ws2["F19"].value)
    assert p["rate"] == pytest.approx(ws2["D24"].value * 100) and p["principal"] == ws2["E24"].value

    ws1 = load_workbook(EXCEL1, data_only=True)["2016-03-28_주택"]
    src = next(r for r in range(13, 103) if ws1.cell(row=r, column=4).value.date().isoformat() == "2026-08-31"
               and (ws1.cell(row=r, column=10).value or 0) > 0)
    assert p["interest"] == ws1.cell(row=src, column=13).value               # 원리금 테이블 값 그대로(수기 보정 없음)
    # 첨부2는 이 칸을 손으로 옮겨 적어 원본과 어긋나 있다(32,361,800 vs 32,361,807) — 자동연결로 해소되는 오차
    assert 0 < p["interest"] - ws2["F24"].value < 10

    # --- 합계·검증·비고·발급 전 경고
    assert d["totals"]["principal"] == d["payments"][0]["principal"] + d["payments"][1]["principal"]
    assert d["totals"]["total"] == d["totals"]["principal"] + d["totals"]["interest"]
    assert all(x["check"] for x in d["balances"])
    assert d["interest_periods"][1]["ipmt"] is True and d["interest_periods"][0]["ipmt"] is False
    assert d["notes"] == ["초록 음영표시한 융자 건의 경우, IPMT함수를 적용해 이자 산출함"]
    assert d["title"].endswith("'26년 3분기 원리금 계산내역")
    assert d["payment_deadline"] == "2026-08-31"
    assert {s["loan_id"] for s in d["sources"]} == {loan_a, loan_b}


def _libreoffice_recalc(xlsx_bytes: bytes):
    soffice = shutil.which("soffice")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in.xlsx"
        src.write_bytes(xlsx_bytes)
        out = Path(tmp) / "out"
        out.mkdir()
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=120, env={"HOME": tmp, "PATH": "/usr/bin:/bin"})
        return load_workbook(io.BytesIO((out / "in.xlsx").read_bytes()), data_only=True)["Sheet1"]


def test_issue_excel_pdf_formulas_and_snapshot_immutability():
    h = admin_headers()
    dorm, loan_a, loan_b = _nazarene_dormitory(h)
    _template(h)
    r = client.post(f"/api/v1/dormitories/{dorm}/notices", json={"due_date": "2026-08-31"}, headers=h)
    assert r.status_code == 201, r.text
    notice = r.json()
    nid = notice["id"]
    snap = notice["generated_data"]

    x = client.get(f"/api/v1/notices/{nid}/download", headers=h)
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(io.BytesIO(x.content))["Sheet1"]
    assert ws["B3"].value == snap["title"]
    assert ws["B9"].value == ACCOUNT["bank_account"] and ws["G9"].value == ACCOUNT["account_holder"]
    assert ws["E13"].value == "=D13-C13"                          # 원본처럼 파생 칸은 수식
    assert ws["G25"].value == "=SUM(G23:G24)"                     # 합계행 위치/수식이 원본(첨부2)과 동일
    assert ws["B25"].value == "합 계"

    if shutil.which("soffice"):
        calc = _libreoffice_recalc(x.content)                      # 수식을 실제로 계산해 값을 검증
        assert calc["E13"].value == 92 and calc["E14"].value == 184
        assert [calc["G18"].value, calc["G19"].value] == ["TRUE", "TRUE"]
        assert calc["E25"].value == snap["totals"]["principal"]
        assert calc["F25"].value == snap["totals"]["interest"]
        assert calc["G25"].value == snap["totals"]["total"]
        pdf = client.get(f"/api/v1/notices/{nid}/download", params={"format": "pdf"}, headers=h)
        assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"

    # 발급 후 원리금 테이블이 바뀌어도(조기상환) 이미 발급한 공문은 그대로 — 스냅샷
    assert client.post(f"/api/v1/loans/{loan_a}/events/early-repayment", json={
        "event_date": "2026-09-30", "repay_type": "PARTIAL", "repay_amount": 100_000_000}, headers=h).status_code == 201
    again = client.get(f"/api/v1/notices/{nid}", headers=h).json()["generated_data"]
    assert again == snap
    x2 = client.get(f"/api/v1/notices/{nid}/download", headers=h)
    assert load_workbook(io.BytesIO(x2.content))["Sheet1"]["G25"].value == "=SUM(G23:G24)"
    # 같은 상환일의 새 미리보기는 달라지지 않는다(상환일이 같고 이벤트가 이후 날짜) — 이후 상환일은 달라짐
    later = client.get(f"/api/v1/dormitories/{dorm}/notices/preview", params={"due_date": "2026-11-30"}, headers=h)
    assert later.status_code == 200
    assert len(client.get(f"/api/v1/dormitories/{dorm}/notices", headers=h).json()) == 1


def test_notice_validation_warnings_and_permissions():
    h = admin_headers()
    spc_a, dorm_a = make_org(h)
    loan_a, _, _ = make_loan(h, dorm=dorm_a)
    spc_h = spc_manager_headers(h, spc_a)
    _, dorm_b = make_org(h)

    # 양식이 없는 상태는 코드가 아니라 데이터 문제 — 안내 후 거부(활성 양식 비활성화로 재현)
    for t in client.get("/api/v1/document-templates", headers=h).json():
        client.post(f"/api/v1/document-templates/{t['id']}/deactivate", headers=h)
    r = client.get(f"/api/v1/dormitories/{dorm_a}/notices/preview", params={"due_date": "2023-05-31"}, headers=h)
    assert r.status_code == 422 and "양식" in r.json()["detail"]

    _template(h)
    ok = client.get(f"/api/v1/dormitories/{dorm_a}/notices/preview", params={"due_date": "2023-05-31"}, headers=h)
    assert ok.status_code == 200 and ok.json()["notes"] == []          # IPMT 건이 없으면 비고 없음
    assert ok.json()["interest_periods"][0]["days"] == 92

    # 청구할 상환액이 없는 날짜 / 상환액 없는 기숙사
    assert client.get(f"/api/v1/dormitories/{dorm_a}/notices/preview", params={"due_date": "2023-06-15"}, headers=h).status_code == 422
    assert client.get(f"/api/v1/dormitories/{dorm_b}/notices/preview", params={"due_date": "2023-05-31"}, headers=h).status_code == 422

    # 납부기한이 토요일이면 자동으로 옮기지 않고 경고(업무기준 미확정)
    sat = client.get(f"/api/v1/dormitories/{dorm_a}/notices/preview",
                     params={"due_date": "2023-05-31", "payment_deadline": "2023-06-03"}, headers=h).json()
    assert sat["payment_deadline"] == "2023-06-03"
    assert any("영업일이 아닙니다" in w and "토요일" in w for w in sat["warnings"])

    # 권한: 양식 등록/발급은 재단 담당자만, SPC 담당자는 자기 기숙사 공문 조회/다운로드만
    assert client.post("/api/v1/document-templates", json={"template_name": "x", "config": ACCOUNT}, headers=spc_h).status_code == 403
    assert client.get(f"/api/v1/dormitories/{dorm_a}/notices/preview", params={"due_date": "2023-05-31"}, headers=spc_h).status_code == 403
    nid = client.post(f"/api/v1/dormitories/{dorm_a}/notices", json={"due_date": "2023-05-31"}, headers=h).json()["id"]
    assert client.post(f"/api/v1/dormitories/{dorm_a}/notices", json={"due_date": "2023-05-31"}, headers=spc_h).status_code == 403
    assert client.get(f"/api/v1/notices/{nid}/download", headers=spc_h).status_code == 200
    other = client.post(f"/api/v1/dormitories/{dorm_b}/notices", json={"due_date": "2023-05-31"}, headers=h)
    assert other.status_code == 422
    assert client.get(f"/api/v1/dormitories/{dorm_b}/notices", headers=spc_h).status_code == 403

    # 발급은 변경이력에 남는다
    logs = client.get(f"/api/v1/loans/{loan_a}/audit-logs", params={"table_name": "generated_document"}, headers=h).json()
    assert len(logs) >= 1 and logs[0]["action"] == "CREATE" and "원리금 계산내역" in logs[0]["after_value"]


def test_same_fund_loans_get_distinguishing_labels():
    h = admin_headers()
    _, dorm = make_org(h)
    l1, _, _ = make_loan(h, dorm=dorm)
    l2, _, _ = make_loan(h, dorm=dorm, amount=50_000_000)
    _template(h)
    d = client.get(f"/api/v1/dormitories/{dorm}/notices/preview", params={"due_date": "2023-05-31"}, headers=h).json()
    labels = [p["label"] for p in d["payments"]]
    assert len(labels) == 2 and all("공공자금관리기금 (" in x for x in labels) and labels[0] != labels[1]
    assert d["totals"]["principal"] == 5_000_000 + 2_500_000
