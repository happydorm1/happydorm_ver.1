"""Phase 19 — 기존 차입금 일괄 이관: 양식, 검사(되돌림), 실행(전부/전무), 기준일 이전 납부완료·예외, 대조표, 이관 취소."""
import io
import uuid
from datetime import date

from openpyxl import load_workbook

from _helpers import admin_headers, client, make_loan, refs

BASE = "/api/v1/admin/migration"


def _template(h):
    r = client.get(f"{BASE}/template", headers=h)
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content))


def _put(ws, rows):
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)


def _book(h, dorms, loans, rates=(), excepts=()):
    wb = _template(h)
    _put(wb["기숙사"], dorms)
    _put(wb["차입금"], loans)
    _put(wb["금리이력"], rates)
    _put(wb["납부예외"], excepts)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _send(h, content, action="validate", cutoff="2024-06-30", mark_paid=True):
    return client.post(f"{BASE}/{action}", headers=h, files={"file": ("기존.xlsx", content, "application/octet-stream")},
                       data={"cutoff_date": cutoff, "mark_paid": str(mark_paid).lower()})


def _loan_row(dorm, code, *, amount=100_000_000, rate="3.0", rate_type="변동", school=None, assigned="2023-02-28",
              first="2023-05-31", months=60, fund="공공자금관리기금", method="원금균등"):
    return [dorm, school, code, fund, amount, assigned, rate_type, "분기", method, None, None, None, first, months, None, rate, None]


def _find_loan(h, code):
    rows = client.get("/api/v1/loans", headers=h).json()
    rows = rows["items"] if isinstance(rows, dict) else rows
    return next((x for x in rows if x["loan_code"] == code), None)


def test_template_has_sheets_lists_and_guide():
    h = admin_headers()
    wb = _template(h)
    assert {"안내", "기숙사", "차입금", "금리이력", "납부예외", "목록"} <= set(wb.sheetnames)
    assert wb["목록"].sheet_state == "hidden"
    heads = [c.value for c in wb["차입금"][1]]
    assert heads[2] == "차입금코드*" and "최초금리(%)*" in heads
    lists = [c.value for c in wb["목록"]["A"]]
    assert "공공자금관리기금" in lists and "주택도시기금" in lists
    assert wb["차입금"].data_validations.dataValidation          # 드롭다운


def test_validate_reports_row_errors_and_writes_nothing():
    h = admin_headers()
    run = uuid.uuid4().hex[:6]
    existing, _, _ = make_loan(h)
    existing_code = client.get(f"/api/v1/loans/{existing}", headers=h).json()["loan_code"]
    content = _book(h, dorms=[[f"이관D-{run}", "검사대", f"이관S-{run}", None, None],
                              [f"이관D-{run}", "검사대", f"이관S-{run}", None, None]],       # 같은 기숙사 두 번
                    loans=[_loan_row(f"이관D-{run}", existing_code),                      # 이미 있는 코드
                           _loan_row(f"이관D-{run}", f"M-{run}-2", fund="없는재원"),
                           _loan_row(f"이관D-{run}", f"M-{run}-3", rate="abc", assigned="2019/13/40"),
                           [f"이관D-{run}", None, f"M-{run}-4", None, 1, "2023-01-01"]],     # 필수 칸 빔
                    excepts=[[f"M-{run}-9", "2023-05-31", 0, 0, None]])                   # 시트에 없는 코드
    r = _send(h, content)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["ok"] is False and body["applied"] is False
    msgs = [(e["sheet"], e["row"], e["message"]) for e in body["errors"]]
    text = str(msgs)
    assert "두 번" in text and "이미 등록된 차입금코드" in text and "없는 재원" in text and "금리가 숫자가 아닙니다" in text and "없는 날짜" in text
    assert "필수 칸이 비어" in text and "차입금' 시트에 없는" in text
    assert any(s == "차입금" and row == 3 for s, row, _ in msgs)            # 행 번호(엑셀 기준)
    # 실행도 막힌다 — 아무것도 생기지 않음
    assert _send(h, content, "apply").json()["applied"] is False
    assert _find_loan(h, f"M-{run}-2") is None
    assert not [d for d in client.get("/api/v1/dormitories", headers=h).json() if d["dormitory_name"] == f"이관D-{run}"]


def test_apply_creates_everything_marks_paid_with_exceptions_report_and_cancel():
    h = admin_headers()
    run = uuid.uuid4().hex[:6]
    dorm = f"이관D-{run}"
    fixed, var = f"M-{run}-F", f"M-{run}-V"
    content = _book(
        h,
        dorms=[[dorm, "이관대학교", f"이관S-{run}", "행복기숙사(사립)", None]],
        loans=[_loan_row(dorm, fixed, rate_type="고정", rate=2.5),
               _loan_row(dorm, var, rate="3.0%", school="이관대학교")],
        rates=[[var, date(2023, 9, 1), 3.5], [var, "2024-03-01", "4.0"]],
        excepts=[[var, "2023-11-30", 0, 0, None],                       # 미납
                 [var, "2024-02-29", 1_000_000, 0, "2024-03-05"]])      # 일부납부
    v = _send(h, content)
    assert v.json()["ok"] is True, v.json()["errors"]
    assert _find_loan(h, fixed) is None                                  # 검사는 되돌린다

    r = _send(h, content, "apply").json()
    assert r["applied"] is True and r["batch_id"], r
    st = r["stats"]
    assert st["dorm_new"] == 1 and st["spc_new"] == 1 and st["loans"] == 2 and st["rates"] == 2
    assert st["unpaid_exceptions"] == 1 and st["payments_exception"] == 1

    # 고정금리: 같은 조건으로 화면에서 만든 차입금과 원리금이 같다
    f = _find_loan(h, fixed)
    sched_f = client.get(f"/api/v1/loans/{f['id']}/schedule", headers=h).json()
    ref_id, _, _ = make_loan(h, rate="2.5")
    ref = client.get(f"/api/v1/loans/{ref_id}/schedule", headers=h).json()
    pick = lambda s: [(x["due_date"], x["principal_amount"], x["interest_amount"]) for x in s["rows"]]  # noqa: E731
    assert pick(sched_f) == pick(ref)

    # 변동금리: 금리이력 반영(이자계산종료일 기준)
    vl = _find_loan(h, var)
    rows = client.get(f"/api/v1/loans/{vl['id']}/schedule", headers=h).json()["rows"]
    rate_at = {x["due_date"]: float(x["applied_rate"]) for x in rows}
    assert rate_at["2023-05-31"] == 3.0 and rate_at["2023-11-30"] == 3.5 and rate_at["2024-05-31"] == 4.0
    hist = client.get(f"/api/v1/loans/{vl['id']}", headers=h).json()
    assert hist["rate_type"] == "VARIABLE"

    # 납부: 기준일(2024-06-30)까지 회차는 예정대로, 예외는 그대로 → 미납 1건·일부납부 1건만 남음
    summ = client.get(f"/api/v1/loans/{vl['id']}/payments/summary", params={"as_of": "2026-10-28"}, headers=h).json()
    status = {x["due_date"]: x["status"] for x in summ["rows"]}
    past = [d for d in status if d <= "2024-06-30"]
    assert status["2023-11-30"] == "UNPAID" and status["2024-02-29"] == "PARTIAL"
    assert all(status[d] == "PAID" for d in past if d not in ("2023-11-30", "2024-02-29"))
    nxt = min(d for d in status if d > "2024-06-30")
    assert status[nxt] == "UNPAID"                                          # 기준일 뒤 회차는 손대지 않음
    fs = client.get(f"/api/v1/loans/{f['id']}/payments/summary", params={"as_of": "2024-06-30"}, headers=h).json()
    assert all(x["status"] == "PAID" for x in fs["rows"] if x["due_date"] <= "2024-06-30")
    pays = client.get(f"/api/v1/loans/{f['id']}/payments", headers=h).json()
    assert pays and "이관" in (pays[0].get("note") or pays[0].get("remark") or "")

    # 같은 파일 다시 실행 → 코드 중복으로 막힘
    again = _send(h, content, "apply").json()
    assert again["applied"] is False and "이미 등록된 차입금코드" in str(again["errors"])

    # 이력·대조표
    b = next(x for x in client.get(f"{BASE}/batches", headers=h).json() if x["id"] == r["batch_id"])
    assert b["cutoff_date"] == "2024-06-30" and len(b["loan_ids"]) == 2 and b["created_by_name"]
    rep = client.get(f"{BASE}/batches/{b['id']}/report", headers=h)
    assert rep.status_code == 200
    ws = load_workbook(io.BytesIO(rep.content))["이관 대조표"]
    head = [c.value for c in ws[1]]
    assert head[0] == "차입금코드" and head[-1] == "차이" and "기존 엑셀 잔여원금(직접 입력)" in head
    data = {ws.cell(row=i, column=1).value: [c.value for c in ws[i]] for i in range(2, ws.max_row + 1)}
    vrow = data[var]
    assert vrow[13] == 2 and vrow[14] == "정상"                              # 미납·부족 회차 2, 계산검증 정상
    assert vrow[8] == 100_000_000 - sum(x["principal_amount"] for x in rows if x["due_date"] <= "2024-06-30")
    assert str(vrow[9])[:10] == nxt
    assert str(ws.cell(row=2, column=17).value).startswith("=IF(")

    # 이관 취소 → 차입금·새 기숙사 숨김, 같은 파일 다시 이관 가능
    assert client.post(f"{BASE}/batches/{b['id']}/cancel", json={"reason": ""}, headers=h).status_code == 422
    c = client.post(f"{BASE}/batches/{b['id']}/cancel", json={"reason": "E2E 재이관"}, headers=h)
    assert c.status_code == 200 and c.json() == {"loans": 2, "dormitories": 1}
    assert client.post(f"{BASE}/batches/{b['id']}/cancel", json={"reason": "x"}, headers=h).status_code == 409
    assert _find_loan(h, fixed) is None
    dels = str(client.get("/api/v1/admin/deletions", headers=h).json())
    assert "이관 취소" in dels and dorm in dels
    re = _send(h, content, "apply").json()
    assert re["applied"] is True, re["errors"]


def test_existing_dorm_reused_no_mark_paid_and_engine_error_is_row_error():
    h = admin_headers()
    run = uuid.uuid4().hex[:6]
    _, dorm_id, _ = make_loan(h)
    dorm = next(d for d in client.get("/api/v1/dormitories", headers=h).json() if d["id"] == dorm_id)
    code = f"M-{run}-E"
    content = _book(h, dorms=[], loans=[_loan_row(dorm["dormitory_name"], code)])
    r = _send(h, content, "apply", mark_paid=False).json()
    assert r["applied"] is True and r["stats"]["dorm_new"] == 0 and r["stats"]["payments_full"] == 0
    loan = _find_loan(h, code)
    assert loan["dormitory_id"] == dorm_id
    assert client.get(f"/api/v1/loans/{loan['id']}/payments", headers=h).json() == []

    # 금리 적용시작일이 배정일보다 늦어 첫 회차 금리가 없음 → 그 행의 오류로 보고(500 아님)
    bad = _loan_row(dorm["dormitory_name"], f"M-{run}-X")
    bad[16] = "2024-01-01"
    r2 = _send(h, _book(h, dorms=[], loans=[bad])).json()
    assert r2["ok"] is False and "원리금 테이블을 만들지 못했습니다" in str(r2["errors"])
    assert r2["errors"][0]["row"] == 2

    # 미래 기준일·엑셀이 아닌 파일
    assert "오늘 이후" in str(_send(h, content, cutoff="2999-01-01").json()["errors"])
    assert "읽지 못했습니다" in str(_send(h, b"not excel").json()["errors"])


def test_non_admin_forbidden():
    h = admin_headers()
    email = f"fund-{uuid.uuid4().hex[:6]}@example.com"
    assert client.post("/api/v1/users", json={"name": "자금", "email": email, "role": "FOUNDATION_FUND_MANAGER",
                                               "password": "fundpass123"}, headers=h).status_code == 201
    tok = client.post("/api/v1/auth/login", data={"username": email, "password": "fundpass123"}).json()["access_token"]
    fh = {"Authorization": f"Bearer {tok}"}
    assert client.get(f"{BASE}/template", headers=fh).status_code == 403
    assert client.get(f"{BASE}/batches", headers=fh).status_code == 403
    assert _send(fh, b"x").status_code == 403
    refs(h)
