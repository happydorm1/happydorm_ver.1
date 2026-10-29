"""Phase 15 최종 검증: §37 계산검증 9개 항목, §50 전체 업무 흐름(기숙사 등록 → … → 대시보드 자동 반영)을 한 번에 통과하는지"""
from datetime import date

from _helpers import admin_headers, client, make_loan, make_org, refs, spc_manager_headers
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.schedule import RepaymentScheduleRow, RepaymentScheduleVersion


def _verify(h, loan, **p):
    r = client.get(f"/api/v1/loans/{loan}/schedule/verification", params=p, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _status(v):
    return {c["no"]: c["status"] for c in v["checks"]}


def test_verification_nine_checks_pass_for_every_event_type():
    h = admin_headers()
    # 원금균등 + 금리변동
    a, _, _ = make_loan(h)
    v = _verify(h, a)
    assert v["overall"] == "NORMAL" and len(v["checks"]) == 9
    assert _status(v)[9] == "NA"                                                      # 최초 버전
    assert client.post(f"/api/v1/loans/{a}/events/rate-change", json={"effective_start_date": "2024-08-31", "annual_rate": "3.4"},
                       headers=h).status_code == 201
    v = _verify(h, a)
    assert v["overall"] == "NORMAL" and _status(v)[9] == "PASS" and "이전 회차 변경 없음" in v["checks"][8]["detail"]
    assert _verify(h, a, version_no=1)["checks"][6]["status"] == "NA"                 # 지난 버전은 금리 비교 제외

    # 원리금균등 + 상환유예(회수원금이 잔액 계산에 들어가는지)
    b, _, _ = make_loan(h, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT", first="2023-08-31", rate="2.0", months=120)
    assert client.post(f"/api/v1/loans/{b}/events/deferral", json={"event_date": "2024-02-29", "defer_target": "PRINCIPAL_INTEREST",
                                                                  "recovery_cycle_months": 3, "recovery_count": 4}, headers=h).status_code == 201
    v = _verify(h, b)
    assert v["overall"] == "NORMAL", [c for c in v["checks"] if c["status"] == "FAIL"]
    assert "상환유예 회수원금" in v["checks"][4]["detail"]

    # 가속상환·조기상환
    c, _, _ = make_loan(h, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT", first="2023-08-31", rate="2.0", months=120)
    client.post(f"/api/v1/loans/{c}/events/acceleration", json={"event_date": "2024-01-15"}, headers=h).raise_for_status()
    assert _verify(h, c)["overall"] == "NORMAL"
    d, _, _ = make_loan(h)
    client.post(f"/api/v1/loans/{d}/events/early-repayment", json={"event_date": "2023-11-30", "repay_type": "PARTIAL",
                                                                   "repay_amount": 10_000_000}, headers=h).raise_for_status()
    assert _verify(h, d)["overall"] == "NORMAL"


def test_verification_detects_tampering_and_unapplied_rate():
    h = admin_headers()
    loan, _, _ = make_loan(h)
    db = SessionLocal()
    v = db.execute(select(RepaymentScheduleVersion).where(RepaymentScheduleVersion.loan_id == loan,
                                                          RepaymentScheduleVersion.is_current.is_(True))).scalar_one()
    row = db.execute(select(RepaymentScheduleRow).where(RepaymentScheduleRow.version_id == v.id, RepaymentScheduleRow.installment_no == 3)).scalar_one()
    row.principal_amount += 1          # DB를 직접 고쳐 계산 근거를 깨뜨림
    row.days += 1
    db.commit()
    db.close()
    res = _verify(h, loan)
    assert res["overall"] == "NEEDS_REVIEW" and set(res["failed"]) >= {1, 5, 6}

    other, _, _ = make_loan(h)
    client.post(f"/api/v1/loans/{other}/rate-history", json={"effective_start_date": "2024-08-31", "annual_rate": "3.6"}, headers=h)
    res = _verify(h, other)
    assert res["failed"] == [7] and "3.6" in res["checks"][6]["items"][0]


def test_end_to_end_business_flow_section_50():
    """§50: 기숙사 등록 → 차입금 → 원리금 테이블 → 검증 → 버전 → 이벤트 → 새 버전 → 알림 → 공문 → 납부 → 비교 → 이력 → 대시보드"""
    h = admin_headers()
    funds, cyc, mth = refs(h)
    spc, dorm = make_org(h)                                                          # 기숙사 등록
    loan = client.post("/api/v1/loans", json={                                        # 차입금 등록 + 기본조건
        "dormitory_id": dorm, "fund_id": funds["공공자금관리기금"]["id"], "loan_code": f"E2E-{dorm[:6]}", "assigned_amount": 200_000_000,
        "assigned_date": "2024-02-29", "rate_type": "VARIABLE", "repay_cycle_code_id": cyc["QUARTERLY"]["id"],
        "repay_method_code_id": mth["EQUAL_PRINCIPAL"]["id"], "grace_start_date": "2024-02-29", "grace_end_date": "2025-02-28",
        "principal_repay_start_date": "2025-05-31", "repay_period_months": 120,
        "initial_rate": {"effective_start_date": "2024-02-29", "effective_end_date": "2025-11-30", "annual_rate": "2.5"}}, headers=h).json()["id"]
    v1 = client.post(f"/api/v1/loans/{loan}/schedule/generate", headers=h).json()          # 원리금 계산·테이블 생성
    assert v1["status"] == "NORMAL" and v1["version_no"] == 1
    assert _verify(h, loan)["overall"] == "NORMAL"                                         # 검증

    as_of = "2025-11-20"
    dash = client.get("/api/v1/dashboard", params={"as_of": as_of, "dormitory_id": dorm}, headers=h).json()
    assert dash["alerts"]["rate_changes_upcoming"]["count"] == 1                            # 금리변동 사전알림(영업일 D-30)
    upcoming = dash["alerts"]["repayment_upcoming"]["items"]
    assert upcoming and upcoming[0]["due_date"] == "2025-11-30"                            # 상환일 도래 사전알림
    scan = client.post("/api/v1/notifications/scan", params={"as_of": as_of, "dormitory_id": dorm}, headers=h).json()
    assert scan["created"]["rate_change"] >= 1

    # 이벤트: 금리변동 → 새 버전(지난 회차 불변)
    v2 = client.post(f"/api/v1/loans/{loan}/events/rate-change", json={"effective_start_date": "2025-11-30", "annual_rate": "2.9"},
                     headers=h).json()
    assert v2["version_no"] == 2
    ver = _verify(h, loan)
    assert ver["overall"] == "NORMAL" and _status(ver)[9] == "PASS"
    assert client.get("/api/v1/dashboard", params={"as_of": as_of, "dormitory_id": dorm},
                      headers=h).json()["alerts"]["needs_attention"]["midterm_pending"] == []

    # 안내공문 자동생성(새 버전 값 연결)
    tpl = client.post("/api/v1/document-templates", json={"template_name": f"E2E-{dorm[:6]}",
                                                         "config": {"bank_account": "검증은행 000-00-000000", "account_holder": "검증"}}, headers=h).json()
    notice = client.post(f"/api/v1/dormitories/{dorm}/notices", json={"template_id": tpl["id"], "due_date": "2025-11-30"}, headers=h).json()
    due_row = [r for r in v2["rows"] if r["due_date"] == "2025-11-30"]
    assert notice["generated_data"]["totals"]["total"] == sum(r["principal_amount"] + r["interest_amount"] + r["unpaid_interest_amount"] for r in due_row)

    # 실제 납부(일부) → 예정/실제 비교 → 이력 → 대시보드 반영
    sched = sum(r["principal_amount"] + r["interest_amount"] for r in due_row)
    pay = client.post(f"/api/v1/loans/{loan}/payments", json={"due_date": "2025-11-30", "actual_paid_date": "2025-11-30",
                                                              "actual_principal": due_row[-1]["principal_amount"], "actual_interest": 0}, headers=h)
    assert pay.status_code == 201
    summ = client.get(f"/api/v1/loans/{loan}/payments/summary", params={"as_of": "2025-12-15"}, headers=h).json()
    row = next(r for r in summ["rows"] if r["due_date"] == "2025-11-30")
    assert row["status"] == "PARTIAL" and row["outstanding"] == sched - due_row[-1]["principal_amount"]
    later = client.get("/api/v1/dashboard", params={"as_of": "2025-12-15", "dormitory_id": dorm}, headers=h).json()
    item = next(x for x in later["alerts"]["overdue"]["items"] if x["due_date"] == "2025-11-30")
    assert item["status"] == "PARTIAL" and item["outstanding"] == row["outstanding"]          # 일부납부가 미납 알림·KPI로
    assert later["kpis"]["overdue_count"] == later["alerts"]["overdue"]["count"]

    hist = client.get(f"/api/v1/loans/{loan}/audit-logs", headers=h).json()
    tables = {x["table_name"] for x in (hist if isinstance(hist, list) else hist.get("items", []))}
    assert {"loan", "interest_rate_history", "loan_event", "actual_payment", "generated_document"} <= tables   # 모든 변경 추적
    versions = client.get(f"/api/v1/loans/{loan}/schedule/versions", headers=h).json()
    assert [x["version_no"] for x in versions] == [1, 2] and versions[0]["is_current"] is False          # 덮어쓰지 않고 보존
    assert date.fromisoformat(as_of) < date(2025, 11, 30)


def test_schedule_rows_and_excel_show_deferral_recovery_principal():
    """상환유예 회수일의 원금상환액에 회수원금이 포함돼 '상환 전 잔액 - 원금상환액 = 상환 후 잔액'이 화면·Excel에서 그대로 성립"""
    import io

    from openpyxl import load_workbook
    h = admin_headers()
    loan, _, _ = make_loan(h, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT", first="2023-08-31", rate="2.0", months=120)
    client.post(f"/api/v1/loans/{loan}/events/deferral", json={"event_date": "2024-02-29", "defer_target": "PRINCIPAL_INTEREST",
                                                               "recovery_cycle_months": 3, "recovery_count": 4}, headers=h).raise_for_status()
    rows = client.get(f"/api/v1/loans/{loan}/schedule", headers=h).json()["rows"]
    rec = [r for r in rows if r["deferral_recovery_principal"]]
    assert len(rec) == 4
    for r in rows:
        if r["row_type"] != "YEAR_END_SPLIT_MEMO":
            assert r["balance_before"] - r["principal_total"] == r["balance_after"], r["due_date"]
    ws = load_workbook(io.BytesIO(client.get(f"/api/v1/loans/{loan}/schedule/download", headers=h).content)).active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[16] != "YEAR_END_SPLIT_MEMO":
            assert row[7] - row[9] == row[14]                    # 상환 전 원금 - 원금상환액 = 상환 후 잔액


def test_excel_reports_eight_kinds_scope_and_permissions():
    """§30 Excel 다운로드 8종 — 값은 원리금 테이블·납부 현황과 같고, SPC 담당자는 자기 것만, 다운로드 제한 조회자는 막힘"""
    import io
    import uuid

    from openpyxl import load_workbook
    h = admin_headers()
    funds, _, _ = refs(h)
    spc, dorm = make_org(h)
    a, _, _ = make_loan(h, dorm=dorm)
    b, _, _ = make_loan(h, dorm=dorm, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT", first="2023-08-31", rate="2.0")
    _, other = make_org(h)
    make_loan(h, dorm=other)
    client.post(f"/api/v1/loans/{a}/events/rate-change", json={"effective_start_date": "2024-08-31", "annual_rate": "3.3"}, headers=h)
    pay = client.get(f"/api/v1/loans/{a}/payments/summary", params={"as_of": "2024-06-15"}, headers=h).json()["rows"][0]
    client.post(f"/api/v1/loans/{a}/payments", json={"due_date": pay["due_date"], "actual_paid_date": pay["due_date"],
                                                     "actual_principal": pay["scheduled_principal"],
                                                     "actual_interest": pay["scheduled_interest"] + 1}, headers=h).raise_for_status()

    def get(kind, hdr=h, **p):
        r = client.get("/api/v1/reports/excel", params={"kind": kind, "as_of": "2024-06-15", **p}, headers=hdr)
        assert r.status_code == 200, (kind, r.text)
        return load_workbook(io.BytesIO(r.content))

    wb = get("schedules_dormitory", dormitory_id=dorm)
    assert wb.sheetnames[0] == "목록" and len(wb.sheetnames) == 3
    sheet = wb[wb.sheetnames[1]]
    vals = [r for r in sheet.iter_rows(min_row=2, values_only=True)]
    assert sum(r[9] for r in vals if r[16] != "YEAR_END_SPLIT_MEMO") in (100_000_000,)          # 원금상환액 합 = 배정금액
    assert len(get("schedules_fund", fund_id=funds["주택도시기금"]["id"]).sheetnames) >= 2
    per = get("period_repayments", dormitory_id=dorm, date_from="2024-01-01", date_to="2024-12-31").active
    body = list(per.iter_rows(min_row=2, values_only=True))
    total = body[-1]
    assert total[0] == "합계" and total[7] == sum(r[7] for r in body[:-1])
    ap = list(get("actual_payments", dormitory_id=dorm).active.iter_rows(min_row=2, values_only=True))
    paid = next(r for r in ap if r[1] == next(x for x in client.get("/api/v1/loans", headers=h).json() if x["id"] == a)["loan_code"]
                and str(r[3])[:10] == pay["due_date"])
    assert paid[14] == "초과납부" and paid[12] == 1
    ev = list(get("midterm_events", dormitory_id=dorm).active.iter_rows(min_row=2, values_only=True))
    assert any(r[2] == "금리변동" and r[4] == "적용" for r in ev)
    rh = list(get("rate_history", dormitory_id=dorm).active.iter_rows(min_row=2, values_only=True))
    assert {r[5] for r in rh} >= {3.0, 3.3, 2.0}
    nd = list(get("notice_data", dormitory_id=dorm, due_date=pay["due_date"]).active.iter_rows(min_row=2, values_only=True))
    assert len(nd) >= 1 and all(r[13] == r[9] + r[12] for r in nd)

    spc_h = spc_manager_headers(h, spc)
    wb = get("schedules_all", hdr=spc_h)
    assert len(wb.sheetnames) == 3                                                              # 자기 기숙사 2건만
    email = f"v-{uuid.uuid4().hex[:6]}@example.com"
    client.post("/api/v1/users", json={"name": "v", "email": email, "role": "VIEWER", "password": "Secure!2345", "can_download": False}, headers=h)
    vh = {"Authorization": "Bearer " + client.post("/api/v1/auth/login", data={"username": email, "password": "Secure!2345"}).json()["access_token"]}
    assert client.get("/api/v1/reports/excel", params={"kind": "rate_history"}, headers=vh).status_code == 403
    assert client.get("/api/v1/reports/excel", params={"kind": "nope"}, headers=h).status_code == 422
    assert client.get("/api/v1/reports/excel", params={"kind": "period_repayments"}, headers=h).status_code == 422
