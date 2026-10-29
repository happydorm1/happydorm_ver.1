import sys
import uuid
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.main import app  # noqa: E402
from app.models.organization import User, UserRole  # noqa: E402

client = TestClient(app)
EXCEL4 = "/mnt/user-data/uploads/첨부_4__원금가속상환_사례_테이블.xlsx"


def _headers():
    db = SessionLocal()
    email = "pytest-admin@example.com"
    if not db.query(User).filter(User.email == email).first():
        db.add(User(name="테스트관리자", email=email, role=UserRole.SYSTEM_ADMIN,
                     hashed_password=hash_password("testpass123")))
        db.commit()
    db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": "testpass123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _refs(h):
    funds = {f["fund_name"]: f for f in client.get("/api/v1/reference/funds", headers=h).json()}
    cyc = {c["code"]: c for c in client.get("/api/v1/reference/codes", params={"group_code": "REPAY_CYCLE"}, headers=h).json()}
    mth = {c["code"]: c for c in client.get("/api/v1/reference/codes", params={"group_code": "REPAY_METHOD"}, headers=h).json()}
    return funds, cyc, mth


def _make_loan(h, *, fund, cycle, method, start, first, months, amount=100_000_000, rate="3.00000",
               grace=None, rate_type="FIXED"):
    funds, cyc, mth = _refs(h)
    run = uuid.uuid4().hex[:8]
    spc = client.post("/api/v1/spc", json={"spc_name": f"S-{run}"}, headers=h).json()["id"]
    dorm = client.post("/api/v1/dormitories", json={"dormitory_name": f"D-{run}", "school_name": "t", "spc_id": spc}, headers=h).json()["id"]
    payload = {
        "dormitory_id": dorm, "fund_id": funds[fund]["id"], "loan_code": f"P8-{run}",
        "assigned_amount": amount, "assigned_date": start, "rate_type": rate_type,
        "repay_cycle_code_id": cyc[cycle]["id"], "repay_method_code_id": mth[method]["id"],
        "principal_repay_start_date": first, "repay_period_months": months,
        "initial_rate": {"effective_start_date": start, "annual_rate": rate},
    }
    if grace:
        payload["grace_start_date"], payload["grace_end_date"] = grace
    r = client.post("/api/v1/loans", json=payload, headers=h)
    assert r.status_code == 201, r.text
    loan_id = r.json()["id"]
    assert client.post(f"/api/v1/loans/{loan_id}/schedule/generate", headers=h).status_code == 201
    return loan_id


# ---------------- 원금가속상환 ----------------
def test_acceleration_matches_attachment_4():
    if not Path(EXCEL4).exists():
        pytest.skip("첨부4 파일이 없어 스킵")
    import openpyxl

    h = _headers()
    loan_id = _make_loan(
        h, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT",
        start="2019-12-30", first="2023-02-28", months=324, amount=15_055_000_000, rate="1.57200",
        grace=("2019-12-30", "2022-08-31"),
    )
    r = client.post(f"/api/v1/loans/{loan_id}/events/acceleration",
                    json={"event_date": "2026-11-30", "acceleration_multiplier": "1.5"}, headers=h)
    assert r.status_code == 201, r.text
    v2 = r.json()
    assert v2["status"] == "NORMAL"

    ws = openpyxl.load_workbook(EXCEL4, data_only=True)["2019-12-30_주택(가속상환)"]
    act = []
    for rr in range(13, 200):
        if ws.cell(row=rr, column=2).value is None:
            break
        act.append({"bal": ws.cell(row=rr, column=9).value, "K": ws.cell(row=rr, column=11).value or 0,
                    "L": ws.cell(row=rr, column=12).value or 0, "M": ws.cell(row=rr, column=13).value or 0})
    assert len(v2["rows"]) == len(act) == 71

    for i, (g, a) in enumerate(zip(v2["rows"], act)):
        if i < 33:  # 가속 시작 이전 구간(원본과 동일, V1 검증)
            continue
        expected_interest = a["M"] if g["installment_no"] is None else a["L"]
        assert g["principal_amount"] == a["K"], f"행{i+1} 원금"
        assert g["balance_before"] == a["bal"], f"행{i+1} 잔액"
        assert g["interest_amount"] == expected_interest, f"행{i+1} 이자"

    assert v2["rows"][-1]["installment_no"] == 41
    assert v2["rows"][-1]["balance_after"] == 0


def test_acceleration_multiplier_is_variable():
    h = _headers()
    loan_id = _make_loan(h, fund="공공자금관리기금", cycle="QUARTERLY", method="EQUAL_PRINCIPAL",
                         start="2023-02-28", first="2023-05-31", months=60)
    r = client.post(f"/api/v1/loans/{loan_id}/events/acceleration",
                    json={"event_date": "2024-06-01", "acceleration_multiplier": "2.0"}, headers=h)
    assert r.status_code == 201, r.text
    v = r.json()
    assert v["status"] == "NORMAL"
    assert sum(x["principal_amount"] for x in v["rows"]) == 100_000_000
    assert v["rows"][-1]["balance_after"] == 0
    # 가속 후에는 회차 수가 원래(20회)보다 줄어든다
    assert max(x["installment_no"] or 0 for x in v["rows"]) < 20


# ---------------- 조기상환 ----------------
def test_partial_early_repayment_equal_principal_keeps_period():
    h = _headers()
    loan_id = _make_loan(h, fund="공공자금관리기금", cycle="QUARTERLY", method="EQUAL_PRINCIPAL",
                         start="2023-02-28", first="2023-05-31", months=60)
    r = client.post(f"/api/v1/loans/{loan_id}/events/early-repayment",
                    json={"event_date": "2024-02-29", "repay_type": "PARTIAL", "repay_amount": 20_000_000}, headers=h)
    assert r.status_code == 201, r.text
    v = r.json()
    assert v["status"] == "NORMAL"
    after = [x for x in v["rows"] if x["installment_no"] and x["due_date"] > "2024-02-29"]
    assert len(after) == 16  # 남은 회차수 유지(20-4)
    assert {x["principal_amount"] for x in after} == {3_750_000}  # 6천만원 / 16회
    assert sum(x["principal_amount"] for x in v["rows"]) == 100_000_000
    assert v["rows"][-1]["balance_after"] == 0


def test_partial_early_repayment_equal_payment_keeps_period():
    h = _headers()
    loan_id = _make_loan(h, fund="주택도시기금", cycle="SEMIANNUAL", method="EQUAL_PAYMENT",
                         start="2023-02-28", first="2023-08-31", months=60, rate="2.00000")
    r = client.post(f"/api/v1/loans/{loan_id}/events/early-repayment",
                    json={"event_date": "2024-08-31", "repay_type": "PARTIAL",
                          "repay_amount": 30_000_000, "fee_amount": 150_000}, headers=h)
    assert r.status_code == 201, r.text
    v = r.json()
    assert v["status"] == "NORMAL"
    after = [x for x in v["rows"] if x["installment_no"] and x["due_date"] > "2024-08-31"]
    assert len(after) == 7
    assert sum(x["principal_amount"] for x in v["rows"]) == 100_000_000
    assert v["rows"][-1]["balance_after"] == 0


def test_full_early_repayment_settles_accrued_interest_and_ends_schedule():
    h = _headers()
    loan_id = _make_loan(h, fund="공공자금관리기금", cycle="QUARTERLY", method="EQUAL_PRINCIPAL",
                         start="2023-02-28", first="2023-05-31", months=60)
    r = client.post(f"/api/v1/loans/{loan_id}/events/early-repayment",
                    json={"event_date": "2024-04-15", "repay_type": "FULL"}, headers=h)
    assert r.status_code == 201, r.text
    v = r.json()
    last = v["rows"][-1]
    days = (date(2024, 4, 15) - date(2024, 2, 29)).days
    assert last["due_date"] == "2024-04-15"
    assert last["principal_amount"] == 80_000_000          # 4회차 상환 후 잔액 전액
    assert last["interest_amount"] == round(80_000_000 * days * 3.0 / 100 / 365)  # 경과이자 정산
    assert last["balance_after"] == 0
    assert all(x["due_date"] <= "2024-04-15" for x in v["rows"])  # 이후 회차 종료 처리


def test_early_repayment_validation_errors():
    h = _headers()
    loan_id = _make_loan(h, fund="공공자금관리기금", cycle="QUARTERLY", method="EQUAL_PRINCIPAL",
                         start="2023-02-28", first="2023-05-31", months=60)
    # 잔여원금 이상 금액은 PARTIAL로 처리 불가
    r = client.post(f"/api/v1/loans/{loan_id}/events/early-repayment",
                    json={"event_date": "2024-02-29", "repay_type": "PARTIAL", "repay_amount": 999_000_000}, headers=h)
    assert r.status_code == 422
