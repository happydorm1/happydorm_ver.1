import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.main import app  # noqa: E402
from app.models.organization import User, UserRole  # noqa: E402

client = TestClient(app)


def _admin_headers():
    db = SessionLocal()
    email = "pytest-admin@example.com"
    if not db.query(User).filter(User.email == email).first():
        db.add(User(name="테스트관리자", email=email, role=UserRole.SYSTEM_ADMIN,
                     hashed_password=hash_password("testpass123")))
        db.commit()
    db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": "testpass123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_version_list_and_excel_download():
    headers = _admin_headers()
    run_id = uuid.uuid4().hex[:8]

    r = client.post("/api/v1/spc", json={"spc_name": f"SPC-{run_id}"}, headers=headers)
    spc_id = r.json()["id"]
    r = client.post(
        "/api/v1/dormitories",
        json={"dormitory_name": f"기숙사-{run_id}", "school_name": "테스트대", "spc_id": spc_id},
        headers=headers,
    )
    dorm_id = r.json()["id"]

    funds = client.get("/api/v1/reference/funds", headers=headers).json()
    fund = next(f for f in funds if f["fund_name"] == "공공자금관리기금")
    cycles = client.get("/api/v1/reference/codes", params={"group_code": "REPAY_CYCLE"}, headers=headers).json()
    quarterly = next(c for c in cycles if c["code"] == "QUARTERLY")
    methods = client.get("/api/v1/reference/codes", params={"group_code": "REPAY_METHOD"}, headers=headers).json()
    equal_principal = next(c for c in methods if c["code"] == "EQUAL_PRINCIPAL")

    r = client.post(
        "/api/v1/loans",
        json={
            "dormitory_id": dorm_id, "fund_id": fund["id"], "loan_code": f"P5-{run_id}",
            "assigned_amount": 100_000_000, "assigned_date": "2024-01-01", "rate_type": "FIXED",
            "repay_cycle_code_id": quarterly["id"], "repay_method_code_id": equal_principal["id"],
            "principal_repay_start_date": "2024-04-30", "repay_period_months": 12,
            "initial_rate": {"effective_start_date": "2024-01-01", "annual_rate": "3.0"},
        },
        headers=headers,
    )
    assert r.status_code == 201, r.text
    loan_id = r.json()["id"]

    r = client.post(f"/api/v1/loans/{loan_id}/schedule/generate", headers=headers)
    assert r.status_code == 201, r.text

    r = client.get(f"/api/v1/loans/{loan_id}/schedule/versions", headers=headers)
    assert r.status_code == 200
    versions = r.json()
    assert len(versions) == 1
    assert versions[0]["version_no"] == 1
    assert versions[0]["generated_reason"] == "최초생성"

    r = client.get(f"/api/v1/loans/{loan_id}/schedule/download", headers=headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]

    import io
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header[0] == "연번"
    assert header[-1] == "비고"
    # 데이터 행 수가 실제 생성된 행 수와 일치
    assert ws.max_row - 1 == len(versions[0]["rows"])
