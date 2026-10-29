"""
Phase 3 (원리금 계산엔진) End-to-End 테스트.
1단계에서 검증한 Test A(공공자금관리기금)/Test B(주택도시기금)를 실제 API + DB로
재현하여, 원본 엑셀(첨부1)과 대조한다.

실행 전: 첨부1 파일이 /mnt/user-data/uploads/첨부_1__원리금_테이블.xlsx 에 있어야 한다.
"""
import sys
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.main import app  # noqa: E402
from app.models.organization import User, UserRole  # noqa: E402

client = TestClient(app)
EXCEL_PATH = "/mnt/user-data/uploads/첨부_1__원리금_테이블.xlsx"


def _admin_headers():
    db = SessionLocal()
    email = "pytest-admin@example.com"
    if not db.query(User).filter(User.email == email).first():
        db.add(User(name="테스트관리자", email=email, role=UserRole.SYSTEM_ADMIN,
                     hashed_password=hash_password("testpass123")))
        db.commit()
    db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": "testpass123"})
    assert r.status_code == 200
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _setup_dormitory(headers, run_id):
    r = client.post("/api/v1/spc", json={"spc_name": f"테스트SPC-{run_id}"}, headers=headers)
    spc_id = r.json()["id"]
    r = client.post(
        "/api/v1/dormitories",
        json={"dormitory_name": f"테스트기숙사-{run_id}", "school_name": "테스트대", "spc_id": spc_id},
        headers=headers,
    )
    return r.json()["id"]


def _code(headers, group_code, code):
    items = client.get("/api/v1/reference/codes", params={"group_code": group_code}, headers=headers).json()
    return next(c for c in items if c["code"] == code)


def _fund(headers, name):
    items = client.get("/api/v1/reference/funds", headers=headers).json()
    return next(f for f in items if f["fund_name"] == name)


def test_test_a_gongja_matches_excel_exactly():
    if not Path(EXCEL_PATH).exists():
        import pytest
        pytest.skip("원본 검증용 엑셀 파일이 없어 스킵합니다.")

    headers = _admin_headers()
    run_id = uuid.uuid4().hex[:8]
    dorm_id = _setup_dormitory(headers, run_id)

    fund = _fund(headers, "공공자금관리기금")
    quarterly = _code(headers, "REPAY_CYCLE", "QUARTERLY")
    equal_principal = _code(headers, "REPAY_METHOD", "EQUAL_PRINCIPAL")

    r = client.post(
        "/api/v1/loans",
        json={
            "dormitory_id": dorm_id, "fund_id": fund["id"], "loan_code": f"TESTA-{run_id}",
            "assigned_amount": 1_768_000_000, "assigned_date": "2015-08-21", "rate_type": "VARIABLE",
            "repay_cycle_code_id": quarterly["id"], "repay_method_code_id": equal_principal["id"],
            "grace_start_date": "2015-08-21", "grace_end_date": "2022-08-31",
            "principal_repay_start_date": "2022-11-30", "repay_period_months": 300,
            "initial_rate": {
                "effective_start_date": "2015-08-21", "effective_end_date": "2022-02-28",
                "annual_rate": "2.24600",
            },
        },
        headers=headers,
    )
    assert r.status_code == 201, r.text
    loan_id = r.json()["id"]

    r = client.post(
        f"/api/v1/loans/{loan_id}/rate-history",
        json={"effective_start_date": "2022-02-28", "annual_rate": "2.32000"},
        headers=headers,
    )
    assert r.status_code == 201

    r = client.post(f"/api/v1/loans/{loan_id}/schedule/generate", headers=headers)
    assert r.status_code == 201, r.text
    version = r.json()
    assert version["status"] == "NORMAL"
    gen_rows = version["rows"]

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["2015-08-21_공자"]
    actual = [
        {
            "J": ws.cell(row=row, column=10).value or 0,
            "K": ws.cell(row=row, column=11).value or 0,
            "L": ws.cell(row=row, column=12).value or 0,
            "days": ws.cell(row=row, column=7).value,
        }
        for row in range(13, 174)
    ]
    assert len(gen_rows) == len(actual) == 161

    mismatches = 0
    for i, (g, a) in enumerate(zip(gen_rows, actual)):
        is_second_of_pair = i > 0 and gen_rows[i - 1]["due_date"] == g["due_date"]
        act_interest = a["K"] if (g["row_type"] != "YEAR_END_SPLIT_REAL" or is_second_of_pair) else a["L"]
        if g["days"] != a["days"] or g["principal_amount"] != a["J"] or g["interest_amount"] != act_interest:
            mismatches += 1

    assert mismatches == 0
    assert sum(r["principal_amount"] for r in gen_rows) == 1_768_000_000
    assert gen_rows[-1]["balance_after"] == 0


def test_test_b_housing_fund_matches_excel_except_known_manual_corrections():
    if not Path(EXCEL_PATH).exists():
        import pytest
        pytest.skip("원본 검증용 엑셀 파일이 없어 스킵합니다.")

    headers = _admin_headers()
    run_id = uuid.uuid4().hex[:8]
    dorm_id = _setup_dormitory(headers, run_id)

    fund = _fund(headers, "주택도시기금")
    semiannual = _code(headers, "REPAY_CYCLE", "SEMIANNUAL")
    equal_payment = _code(headers, "REPAY_METHOD", "EQUAL_PAYMENT")

    r = client.post(
        "/api/v1/loans",
        json={
            "dormitory_id": dorm_id, "fund_id": fund["id"], "loan_code": f"TESTB-{run_id}",
            "assigned_amount": 4_097_000_000, "assigned_date": "2016-03-28", "rate_type": "FIXED",
            "repay_cycle_code_id": semiannual["id"], "repay_method_code_id": equal_payment["id"],
            "grace_start_date": "2016-03-28", "grace_end_date": "2019-02-28",
            "principal_repay_start_date": "2019-08-31", "repay_period_months": 324,
            "initial_rate": {"effective_start_date": "2016-03-28", "annual_rate": "2.00000"},
        },
        headers=headers,
    )
    assert r.status_code == 201, r.text
    loan_id = r.json()["id"]

    r = client.post(f"/api/v1/loans/{loan_id}/schedule/generate", headers=headers)
    assert r.status_code == 201, r.text
    version = r.json()
    assert version["status"] == "NORMAL"
    gen_rows = version["rows"]

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["2016-03-28_주택"]
    actual = [
        {
            "J": ws.cell(row=row, column=10).value or 0,
            "M": ws.cell(row=row, column=13).value or 0,
        }
        for row in range(13, 103)
    ]
    assert len(gen_rows) == len(actual) == 90

    # 52·54회차는 원본 엑셀의 수동보정(-3원/+4원)을 의도적으로 재현하지 않는 유일한 예외(1단계에서 확정)
    KNOWN_EXCEPTION_INSTALLMENTS = {52, 54}
    mismatches = []
    for i, (g, a) in enumerate(zip(gen_rows, actual)):
        if g["principal_amount"] != a["J"] or g["interest_amount"] != a["M"]:
            if g["installment_no"] not in KNOWN_EXCEPTION_INSTALLMENTS:
                mismatches.append((i + 1, g["due_date"], g["installment_no"]))

    assert mismatches == [], f"예상치 못한 불일치: {mismatches}"
    assert sum(r["principal_amount"] for r in gen_rows) == 4_097_000_000
    assert gen_rows[-1]["balance_after"] == 0
