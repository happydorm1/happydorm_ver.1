"""Phase 17: 증빙자료 첨부(암호화 저장·내려받기·형식 제한·권한·사용중지·결재 연결)와 상환일 메모(버전 무관·Excel 반영)"""
import hashlib
import io
import uuid
from pathlib import Path

import pytest
from _helpers import admin_headers, client, make_loan, make_org, spc_manager_headers
from openpyxl import load_workbook

from app.core.config import settings

PW = "Secure!2345"
PDF = b"%PDF-1.4\n% test evidence\n" + b"0123456789" * 200


@pytest.fixture(autouse=True)
def _tmp_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "ATTACHMENT_DIR", str(tmp_path / "attachments"))


def _user(h, role, spc_id=None, **extra):
    email = f"{role.lower()}-{uuid.uuid4().hex[:8]}@example.com"
    body = {"name": role, "email": email, "role": role, "password": PW, **extra}
    if spc_id:
        body["spc_id"] = spc_id
    r = client.post("/api/v1/users", json=body, headers=h)
    assert r.status_code == 201, r.text
    r = client.post("/api/v1/auth/login", data={"username": email, "password": PW})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _upload(h, loan, name="근거.pdf", data=PDF, **form):
    return client.post(f"/api/v1/loans/{loan}/attachments", files={"file": (name, io.BytesIO(data), "application/pdf")},
                       data=form, headers=h)


def test_upload_is_encrypted_at_rest_and_downloads_identical():
    h = admin_headers()
    loan, _, _ = make_loan(h)
    r = _upload(h, loan, description="상환유예 승인 공문")
    assert r.status_code == 201, r.text
    a = r.json()
    assert (a["owner_type"], a["owner_id"], a["size_bytes"]) == ("LOAN", loan, len(PDF))
    assert a["sha256"] == hashlib.sha256(PDF).hexdigest() and a["created_by_name"]
    stored = list(Path(settings.ATTACHMENT_DIR).iterdir())
    assert len(stored) == 1 and PDF[:20] not in stored[0].read_bytes() and stored[0].name != "근거.pdf"   # 원문·원래 이름이 보이지 않게
    d = client.get(f"/api/v1/attachments/{a['id']}/download", headers=h)
    assert d.status_code == 200 and d.content == PDF
    assert "attachment" in d.headers["content-disposition"] and d.headers["x-content-sha256"] == a["sha256"]
    # 저장 파일이 바뀌면 무결성 검사에서 막힌다
    stored[0].write_bytes(stored[0].read_bytes()[:-1] + b"x")
    assert client.get(f"/api/v1/attachments/{a['id']}/download", headers=h).status_code == 410


def test_file_type_size_and_executable_are_rejected(monkeypatch):
    h = admin_headers()
    loan, _, _ = make_loan(h, generate=False)
    assert _upload(h, loan, name="악성.exe").status_code == 422
    assert _upload(h, loan, name="위장.pdf", data=b"MZ\x90\x00" + b"x" * 100).status_code == 422     # 확장자 위장 실행파일
    assert _upload(h, loan, name="빈파일.pdf", data=b"").status_code == 422
    monkeypatch.setattr(settings, "ATTACHMENT_MAX_MB", 1)
    assert _upload(h, loan, name="큰파일.pdf", data=b"%PDF" + b"0" * (1024 * 1024 + 10)).status_code == 422
    assert _upload(h, loan, name="../../etc/한글 이름.hwp", data=b"HWP Document").json()["file_name"] == "한글 이름.hwp"
    assert client.get("/api/v1/attachments/settings", headers=h).json()["max_mb"] == 1


def test_owner_must_belong_to_loan_and_scope_by_spc():
    h = admin_headers()
    spc, dorm = make_org(h)
    loan, _, _ = make_loan(h, dorm=dorm)
    other, _, _ = make_loan(h)
    assert _upload(h, loan, owner_type="LOAN", owner_id=other).status_code == 422
    ev = client.post(f"/api/v1/loans/{other}/events/early-repayment", json={"event_date": "2024-11-15", "repay_type": "PARTIAL",
                                                                           "repay_amount": 5_000_000}, headers=h)
    ev.raise_for_status()
    other_event = client.get(f"/api/v1/loans/{other}/events", headers=h).json()[0]["id"]
    assert _upload(h, loan, owner_type="LOAN_EVENT", owner_id=other_event).status_code == 422          # 다른 차입금의 이벤트
    assert _upload(h, other, owner_type="LOAN_EVENT", owner_id=other_event).status_code == 201
    assert _upload(h, loan, owner_type="UNKNOWN").status_code == 422
    a = _upload(h, loan).json()
    mine, others = spc_manager_headers(h, spc), spc_manager_headers(h, make_org(h)[0])
    assert len(client.get(f"/api/v1/loans/{loan}/attachments", headers=mine).json()) == 1
    assert client.get(f"/api/v1/attachments/{a['id']}/download", headers=mine).status_code == 200
    assert client.get(f"/api/v1/loans/{loan}/attachments", headers=others).status_code == 403
    assert client.get(f"/api/v1/attachments/{a['id']}/download", headers=others).status_code == 403
    assert _upload(mine, loan).status_code == 403                                                  # SPC 담당자는 올리지 못함


def test_viewer_download_follows_setting_and_access_log():
    h = admin_headers()
    loan, _, _ = make_loan(h, generate=False)
    a = _upload(h, loan).json()
    assert client.get(f"/api/v1/attachments/{a['id']}/download", headers=_user(h, "VIEWER")).status_code == 200
    assert client.get(f"/api/v1/attachments/{a['id']}/download", headers=_user(h, "VIEWER", can_download=False)).status_code == 403
    logs = client.get("/api/v1/admin/access-logs", params={"action": "DOWNLOAD"}, headers=h).json()
    items = logs if isinstance(logs, list) else logs.get("items", [])
    assert any(f"/attachments/{a['id']}/download" in (x.get("path") or "") for x in items)


def test_deactivate_keeps_record_audits_and_hides_from_default_list():
    h = admin_headers()
    loan, _, _ = make_loan(h, generate=False)
    a = _upload(h, loan).json()
    assert client.post(f"/api/v1/attachments/{a['id']}/deactivate", json={"reason": "x"}, headers=h).status_code == 422  # 사유 2자 이상
    d = client.post(f"/api/v1/attachments/{a['id']}/deactivate", json={"reason": "잘못 올림"}, headers=h).json()
    assert d["active"] is False and d["deactivate_reason"] == "잘못 올림"
    assert client.get(f"/api/v1/loans/{loan}/attachments", headers=h).json() == []
    assert len(client.get(f"/api/v1/loans/{loan}/attachments", params={"include_inactive": "true"}, headers=h).json()) == 1
    assert client.get(f"/api/v1/attachments/{a['id']}/download", headers=h).status_code == 200        # 기록 확인용 내려받기
    audit = client.get(f"/api/v1/loans/{loan}/audit-logs", headers=h).json()
    rows = audit if isinstance(audit, list) else audit.get("items", [])
    acts = {(x.get("table_name"), x.get("action")) for x in rows}
    assert {("attachment", "CREATE"), ("attachment", "CANCEL")} <= acts


def test_approval_links_request_attachments_to_the_new_event(monkeypatch):
    monkeypatch.setattr(settings, "EVENT_APPROVAL_REQUIRED", True)
    h = admin_headers()
    loan, _, _ = make_loan(h)
    fund, approver = _user(h, "FOUNDATION_FUND_MANAGER"), _user(h, "FOUNDATION_APPROVER")
    req = client.post(f"/api/v1/loans/{loan}/events/rate-change", json={"effective_start_date": "2024-08-31", "annual_rate": "3.3"},
                      headers=fund).json()["request"]
    assert _upload(fund, loan, name="금리통지.pdf", owner_type="EVENT_REQUEST", owner_id=req["id"]).status_code == 201
    assert _upload(approver, loan, owner_type="EVENT_REQUEST", owner_id=req["id"]).status_code == 403    # 결재권자는 올리지 않음
    assert len(client.get(f"/api/v1/loans/{loan}/attachments", params={"owner_type": "EVENT_REQUEST", "owner_id": req["id"]},
                          headers=approver).json()) == 1                                                  # 결재 전 확인
    client.post(f"/api/v1/event-requests/{req['id']}/approve", json={"note": "확인"}, headers=approver).raise_for_status()
    ev = client.get(f"/api/v1/loans/{loan}/events", headers=h).json()[0]
    linked = client.get(f"/api/v1/loans/{loan}/attachments", params={"owner_type": "LOAN_EVENT", "owner_id": ev["id"]}, headers=h).json()
    assert [x["file_name"] for x in linked] == ["금리통지.pdf"]
    assert client.get(f"/api/v1/attachments/{linked[0]['id']}/download", headers=h).content == PDF      # 같은 파일을 가리킴
    assert len(list(Path(settings.ATTACHMENT_DIR).iterdir())) == 1                                     # 파일 복사 없음


def test_schedule_notes_survive_new_versions_and_appear_in_excel():
    h = admin_headers()
    spc, dorm = make_org(h)
    loan, _, _ = make_loan(h, dorm=dorm)
    rows = client.get(f"/api/v1/loans/{loan}/schedule", headers=h).json()["rows"]
    due = rows[3]["due_date"]
    assert client.post(f"/api/v1/loans/{loan}/schedule-notes", json={"due_date": "2023-05-30", "note": "x"}, headers=h).status_code == 422
    n = client.post(f"/api/v1/loans/{loan}/schedule-notes", json={"due_date": due, "note": "입금 지연 협의(담당 통화)"}, headers=h)
    assert n.status_code == 201, n.text
    assert n.json()["written_on_version_no"] == 1
    client.post(f"/api/v1/loans/{loan}/schedule-notes", json={"due_date": due, "note": "익일 입금 확인"}, headers=h).raise_for_status()
    assert client.post(f"/api/v1/loans/{loan}/schedule-notes", json={"due_date": due, "note": "x"},
                       headers=spc_manager_headers(h, spc)).status_code == 403
    client.post(f"/api/v1/loans/{loan}/events/rate-change", json={"effective_start_date": "2024-08-31", "annual_rate": "3.3"},
                headers=h).raise_for_status()                                                              # V2
    notes = client.get(f"/api/v1/loans/{loan}/schedule-notes", headers=spc_manager_headers(h, spc)).json()
    assert [x["note"] for x in notes] == ["입금 지연 협의(담당 통화)", "익일 입금 확인"]

    x = client.get(f"/api/v1/loans/{loan}/schedule/download", headers=h)
    ws = load_workbook(io.BytesIO(x.content)).active
    remarks = {str(r[2].value)[:10]: r[16].value for r in ws.iter_rows(min_row=2)}
    assert remarks[due].endswith("메모: 입금 지연 협의(담당 통화) / 익일 입금 확인")

    client.post(f"/api/v1/schedule-notes/{notes[0]['id']}/deactivate", headers=h).raise_for_status()
    assert [x["note"] for x in client.get(f"/api/v1/loans/{loan}/schedule-notes", headers=h).json()] == ["익일 입금 확인"]
    ws = load_workbook(io.BytesIO(client.get(f"/api/v1/loans/{loan}/schedule/download", headers=h).content)).active
    assert {str(r[2].value)[:10]: r[16].value for r in ws.iter_rows(min_row=2)}[due].endswith("메모: 익일 입금 확인")
    # 리포트(원리금 테이블 묶음)에도 반영
    rep = client.get("/api/v1/reports/excel", params={"kind": "schedules_dormitory", "dormitory_id": dorm, "as_of": "2026-10-28"}, headers=h)
    assert rep.status_code == 200, rep.text
    wb = load_workbook(io.BytesIO(rep.content))
    sheet = [s for s in wb.worksheets if s.title != "목록"][0]
    assert any("메모: 익일 입금 확인" in str(r[16].value) for r in sheet.iter_rows(min_row=2))


def test_backup_includes_attachment_archive(tmp_path, monkeypatch):
    import tarfile

    import app.jobs.backup as bk
    monkeypatch.setattr(settings, "BACKUP_DIR", str(tmp_path / "bk"))
    h = admin_headers()
    loan, _, _ = make_loan(h, generate=False)
    _upload(h, loan)
    log = bk.run_backup()
    assert log.status == "SUCCESS", log.message
    arc = bk.files_archive(Path(settings.BACKUP_DIR) / log.file_name)
    assert arc.exists() and "증빙파일" in (log.message or "")
    with tarfile.open(arc) as tf:
        assert len(tf.getnames()) == 1
