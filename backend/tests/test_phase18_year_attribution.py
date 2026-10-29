"""Phase 18 — 회계연도(1/1~12/31) 기간귀속 이자: 원리금균등 연도경계 회차 이자를 일수 비율로 안분."""
import io
import uuid
from collections import defaultdict

from openpyxl import load_workbook

from _helpers import admin_headers, client, make_org, refs
from app.engine.schedule_engine import excel_round


def _assert_continuous(rows, start):
    """이자계산기간이 빠짐없이 이어진다: 각 행의 시작일 = 앞 행의 끝일."""
    assert rows[0]["interest_calc_start"] == start
    for a, b in zip(rows, rows[1:]):
        assert b["interest_calc_start"] == a["interest_calc_end"], (a["seq_no"], b["seq_no"])


def _make_housing_loan(h):
    """사용자 사례 재현: 주택도시기금 1.572% 반기 원리금균등, 거치 2022-12-23~2025-12-31, 324개월."""
    funds, cyc, mth = refs(h)
    grace = {c["code"]: c for c in client.get("/api/v1/reference/codes", params={"group_code": "GRACE_TYPE"}, headers=h).json()}
    _, dorm = make_org(h)
    body = dict(dormitory_id=dorm, fund_id=funds["주택도시기금"]["id"], loan_code=f"Y-{uuid.uuid4().hex[:6]}",
                assigned_amount=11_512_000_000, assigned_date="2022-12-23", rate_type="FIXED",
                repay_cycle_code_id=cyc["SEMIANNUAL"]["id"], repay_method_code_id=mth["EQUAL_PAYMENT"]["id"],
                grace_type_code_id=grace["GENERAL"]["id"], grace_start_date="2022-12-23", grace_end_date="2025-12-31",
                principal_repay_start_date="2026-02-28", repay_period_months=324,
                initial_rate={"effective_start_date": "2022-12-23", "annual_rate": "1.572"})
    r = client.post("/api/v1/loans", json=body, headers=h)
    assert r.status_code == 201, r.text
    lid = r.json()["id"]
    g = client.post(f"/api/v1/loans/{lid}/schedule/generate", headers=h)
    assert g.status_code == 201, g.text
    return lid, sorted(g.json()["rows"], key=lambda x: x["seq_no"])


def test_year_interest_splits_equal_payment_interest_by_days():
    h = admin_headers()
    lid, rows = _make_housing_loan(h)

    # 1) 총액 보존: 기간귀속 이자 합계 = 원리금 테이블 이자 합계(원단위 오차 없음)
    assert sum(r["year_interest"] for r in rows) == sum(r["interest_amount"] for r in rows)

    # 2) 거치기간(일수 계산)·연도경계 없는 회차는 이자 그대로
    for r in rows:
        if r["row_type"] in ("GRACE", "NORMAL"):
            assert r["year_interest"] == r["interest_amount"]

    # 3) 원리금균등 연도경계: 메모행(0원)에 일수 비율 몫, 납입행에 나머지 — 사용자 사례 2027-02-28 회차
    pair = [r for r in rows if r["due_date"] == "2027-02-28"]
    memo, real = pair
    assert memo["row_type"] == "YEAR_END_SPLIT_MEMO" and memo["interest_amount"] == 0
    assert (memo["days"], real["days"]) == (122, 59) and real["interest_amount"] == 87_770_513
    assert memo["year_interest"] == excel_round(87_770_513 * 122 / 181)
    assert memo["year_interest"] + real["year_interest"] == 87_770_513

    # 4) 거치종료일(12/31)이 납입일 사이여도 기간 누락 없음: 첫 상환회차는 마지막 거치 납입일(2025-08-31)부터
    first = [r for r in rows if r["due_date"] == "2026-02-28"]
    assert [(r["interest_calc_start"], r["interest_calc_end"], r["days"]) for r in first] == [
        ("2025-08-31", "2025-12-31", 122), ("2025-12-31", "2026-02-28", 59)]
    assert first[1]["interest_amount"] == 90_484_320                       # 회차 이자(공식)는 그대로
    assert sum(r["interest_amount"] for r in rows) == 3_146_776_693       # 총 이자도 그대로
    assert not [r for r in rows if r["days"] == 0]
    _assert_continuous(rows, "2022-12-23")

    # 5) Excel: 마지막 열 + '연도별 집계' 시트, 연도별 합계가 API 값과 일치
    x = client.get(f"/api/v1/loans/{lid}/schedule/download", headers=h)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb.worksheets[0]
    assert ws.cell(row=1, column=17).value == "비고"                     # 기존 열 위치 유지
    assert ws.cell(row=1, column=18).value.startswith("기간귀속 이자")
    assert [c.value for c in ws[2]][17] == rows[0]["year_interest"]

    expect = defaultdict(int)
    for r in rows:
        expect[int(r["interest_calc_end"][:4])] += r["year_interest"] + r["unpaid_interest_amount"]
    ys = wb["연도별 집계"]
    got = {row[0]: row[3] for row in ys.iter_rows(min_row=2, values_only=True) if isinstance(row[0], int)}
    assert got == {y: v for y, v in expect.items()}
    assert got[2026] > 0 and got[2027] > 0


def test_equal_principal_grace_end_between_due_dates_has_no_gap():
    """원금균등(공자, 일수 이자)에서도 거치종료일이 분기 납입일 사이면 그 사이 이자가 빠지지 않는다."""
    h = admin_headers()
    funds, cyc, mth = refs(h)
    _, dorm = make_org(h)
    body = dict(dormitory_id=dorm, fund_id=funds["공공자금관리기금"]["id"], loan_code=f"Q-{uuid.uuid4().hex[:6]}",
                assigned_amount=1_200_000_000, assigned_date="2023-01-10", rate_type="FIXED",
                repay_cycle_code_id=cyc["QUARTERLY"]["id"], repay_method_code_id=mth["EQUAL_PRINCIPAL"]["id"],
                grace_start_date="2023-01-10", grace_end_date="2024-01-15",        # 분기 납입일(2/5/8/11월 말) 사이
                principal_repay_start_date="2024-02-29", repay_period_months=120,
                initial_rate={"effective_start_date": "2023-01-10", "annual_rate": "2.0"})
    lid = client.post("/api/v1/loans", json=body, headers=h).json()["id"]
    g = client.post(f"/api/v1/loans/{lid}/schedule/generate", headers=h)
    assert g.status_code == 201, g.text
    rows = sorted(g.json()["rows"], key=lambda x: x["seq_no"])
    _assert_continuous(rows, "2023-01-10")
    first = next(r for r in rows if r["installment_no"] == 1)
    assert first["due_date"] == "2024-02-29" and first["interest_calc_start"] == "2023-12-31"
    gap_row = next(r for r in rows if r["interest_calc_start"] == "2023-11-30")   # 마지막 거치 납입일 이후 기간
    assert gap_row["interest_amount"] > 0
    assert sum(r["year_interest"] for r in rows) == sum(r["interest_amount"] for r in rows)
