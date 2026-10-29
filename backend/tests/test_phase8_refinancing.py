"""
Phase 8 대환대출 테스트 — 첨부6(일부대환)·첨부7(전액대환) 실사례 대조.
두 사례 모두 불균등 상환 원장이라 엔진으로 V1을 만들 수 없으므로, 시트의 '대환 전 구조'를
DB에 직접 적재한 뒤 API로 대환을 적용하여 시트의 대환 이후 행과 비교한다.
"""
import sys
import uuid
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.database import SessionLocal  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.main import app  # noqa: E402
from app.models.organization import User, UserRole  # noqa: E402
from app.models.schedule import (  # noqa: E402
    RepaymentScheduleRow, RepaymentScheduleVersion, RowType, ScheduleVersionStatus,
)

client = TestClient(app)
EXCEL6 = "/mnt/user-data/uploads/첨부_6__일부대환_사례_테이블.xlsx"
EXCEL7 = "/mnt/user-data/uploads/첨부_7__전액대환_사례_테이블.xlsx"


def _headers():
    db = SessionLocal()
    email = "pytest-admin@example.com"
    if not db.query(User).filter(User.email == email).first():
        db.add(User(name="테스트관리자", email=email, role=UserRole.SYSTEM_ADMIN,
                     hashed_password=hash_password("testpass123")))
        db.commit()
    db.close()
    r = client.post("/api/v1/auth/login", data={"username": email, "password": "testpass123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _refs(h):
    funds = {f["fund_name"]: f for f in client.get("/api/v1/reference/funds", headers=h).json()}
    cyc = {c["code"]: c for c in client.get("/api/v1/reference/codes", params={"group_code": "REPAY_CYCLE"}, headers=h).json()}
    mth = {c["code"]: c for c in client.get("/api/v1/reference/codes", params={"group_code": "REPAY_METHOD"}, headers=h).json()}
    return funds, cyc, mth


def _loan_with_ledger(h, ledger, *, assigned, assigned_date, rate, rounding="ROUNDDOWN"):
    """ledger: (due, start, end, days, installment_no, balance_before, principal, interest, is_memo) 목록을 V1으로 직접 적재."""
    funds, cyc, mth = _refs(h)
    run = uuid.uuid4().hex[:6]
    spc = client.post("/api/v1/spc", json={"spc_name": f"S-{run}"}, headers=h).json()["id"]
    dorm = client.post("/api/v1/dormitories", json={"dormitory_name": f"D-{run}", "school_name": "t", "spc_id": spc}, headers=h).json()["id"]
    r = client.post("/api/v1/loans", json={
        "dormitory_id": dorm, "fund_id": funds["공공자금관리기금"]["id"], "loan_code": f"RF-{run}",
        "assigned_amount": assigned, "assigned_date": assigned_date, "rate_type": "VARIABLE",
        "interest_rounding": rounding,
        "repay_cycle_code_id": cyc["QUARTERLY"]["id"], "repay_method_code_id": mth["EQUAL_PRINCIPAL"]["id"],
        "principal_repay_start_date": "2011-03-31", "repay_period_months": 300,
        "initial_rate": {"effective_start_date": assigned_date, "annual_rate": str(rate)},
    }, headers=h)
    assert r.status_code == 201, r.text
    loan_id = r.json()["id"]
    db = SessionLocal()
    v = RepaymentScheduleVersion(loan_id=uuid.UUID(loan_id), version_no=1, generated_reason="최초생성",
                                 is_current=True, status=ScheduleVersionStatus.NORMAL)
    db.add(v)
    db.flush()
    for seq, (due, start, end, days, inst, bb, pr, it, memo) in enumerate(ledger, start=1):
        db.add(RepaymentScheduleRow(
            version_id=v.id, seq_no=seq, due_date=due, interest_calc_start=start, interest_calc_end=end, days=days,
            installment_no=inst, balance_before=bb, principal_amount=pr, interest_amount=it, unpaid_interest_amount=0,
            balance_after=bb - pr, applied_rate=rate,
            row_type=RowType.YEAR_END_SPLIT_MEMO if memo else RowType.NORMAL))
    db.commit()
    db.close()
    return loan_id, funds


def _d(v):
    return v.date() if hasattr(v, "date") else v


def _read_sheet(path, name, last=400):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)[name]
    out = []
    for r in range(13, last):
        if ws.cell(row=r, column=2).value is None:
            break
        out.append({
            "row": r, "fund": ws.cell(row=r, column=3).value, "due": _d(ws.cell(row=r, column=4).value),
            "start": _d(ws.cell(row=r, column=5).value), "end": _d(ws.cell(row=r, column=6).value),
            "days": ws.cell(row=r, column=7).value, "bal": ws.cell(row=r, column=8).value,
            "J": ws.cell(row=r, column=10).value or 0, "K": ws.cell(row=r, column=11).value or 0,
            "L": ws.cell(row=r, column=12).value or 0,
        })
    return out


def _compare(gen_rows, sheet_rows, from_idx):
    bad = []
    for gi, (g, a) in enumerate(zip(gen_rows, sheet_rows)):
        if gi < from_idx:
            continue
        exp_int = a["K"] if a["K"] else a["L"]
        for k, (x, y) in {
            "due": (g["due_date"], a["due"].isoformat()), "days": (g["days"], a["days"]),
            "bal": (g["balance_before"], a["bal"]), "principal": (g["principal_amount"], a["J"]),
            "interest": (g["interest_amount"], exp_int),
        }.items():
            if x != y:
                bad.append((a["row"], k, x, y))
    return bad


def test_partial_refinancing_matches_attachment_6():
    if not Path(EXCEL6).exists():
        pytest.skip("첨부6 파일이 없어 스킵")
    h = _headers()
    S = _read_sheet(EXCEL6, "2010-12-22_공자(일부 대환 예정)")
    win_start = next(i for i, x in enumerate(S) if x["fund"] == "주택도시기금")
    rejoin_idx = next(i for i, x in enumerate(S) if i > win_start and x["fund"] == "공공자금관리기금"
                      and S[i - 1]["fund"] == "공공자금관리기금" and x["due"] >= date(2027, 9, 30))
    before, window, after = S[:win_start], S[win_start:rejoin_idx], S[rejoin_idx:]
    conv = [x for x in window if x["fund"] == "주택도시기금"]
    single = list(before) + [dict(x, fund="합산") for x in conv] + list(after)

    ledger, bal, inst = [], 13_740_000_000, 0
    for i, x in enumerate(single):
        memo = i + 1 < len(single) and single[i + 1]["due"] == x["due"]
        bb = x["bal"] if (x in before or x in after) else bal
        pr = 0 if memo else x["J"]
        if not memo:
            inst += 1
        ledger.append((x["due"], x["start"], x["end"], x["days"], None if memo else inst, bb, pr, x["K"] or x["L"], memo))
        bal = bb - pr

    loan_id, funds = _loan_with_ledger(h, ledger, assigned=13_740_000_000, assigned_date="2010-12-22", rate=4.557)
    r = client.post(f"/api/v1/loans/{loan_id}/events/refinancing", json={
        "refi_type": "PARTIAL", "event_date": "2025-03-31", "new_fund_id": funds["주택도시기금"]["id"],
        "new_annual_rate": "2.17000", "converted_amount": 3_500_000_000,
        "remaining_amount": 4_658_000_000, "rejoin_date": "2027-09-30"}, headers=h)
    assert r.status_code == 201, r.text
    v2 = r.json()
    assert v2["status"] == "NORMAL" and v2["validation_errors"] is None
    assert len(v2["rows"]) == len(S)
    assert _compare(v2["rows"], S, win_start) == []
    assert v2["rows"][-1]["balance_after"] == 0

    # 전환/잔류 트랙이 비고로 구분되고, 재합류 이후는 단일 트랙
    marks = [x["remark"] or "" for x in v2["rows"][win_start:rejoin_idx]]
    assert sum("전환분" in m for m in marks) == len(conv)
    assert sum("잔류분" in m for m in marks) == len(window) - len(conv)


def test_partial_refinancing_rejects_inconsistent_amounts():
    if not Path(EXCEL6).exists():
        pytest.skip("첨부6 파일이 없어 스킵")
    h = _headers()
    S = _read_sheet(EXCEL6, "2010-12-22_공자(일부 대환 예정)")
    win_start = next(i for i, x in enumerate(S) if x["fund"] == "주택도시기금")
    before = S[:win_start]
    conv = [x for x in S[win_start:] if x["fund"] == "주택도시기금"]
    after = [x for x in S[win_start:] if x["fund"] == "공공자금관리기금" and x["due"] >= date(2027, 9, 30)]
    single = list(before) + [dict(x, fund="합산") for x in conv] + list(after)
    ledger, bal, inst = [], 13_740_000_000, 0
    for i, x in enumerate(single):
        memo = i + 1 < len(single) and single[i + 1]["due"] == x["due"]
        bb = x["bal"] if (x in before or x in after) else bal
        pr = 0 if memo else x["J"]
        if not memo:
            inst += 1
        ledger.append((x["due"], x["start"], x["end"], x["days"], None if memo else inst, bb, pr, x["K"] or x["L"], memo))
        bal = bb - pr
    loan_id, funds = _loan_with_ledger(h, ledger, assigned=13_740_000_000, assigned_date="2010-12-22", rate=4.557)

    base = {"refi_type": "PARTIAL", "event_date": "2025-03-31", "new_fund_id": funds["주택도시기금"]["id"],
            "new_annual_rate": "2.17000", "rejoin_date": "2027-09-30"}
    # 전환+잔류 != 잔액
    r = client.post(f"/api/v1/loans/{loan_id}/events/refinancing",
                    json={**base, "converted_amount": 3_500_000_000, "remaining_amount": 4_000_000_000}, headers=h)
    assert r.status_code == 422
    # 합계는 맞지만 전환금액 != 구간 내 정기상환 원금 합계
    r = client.post(f"/api/v1/loans/{loan_id}/events/refinancing",
                    json={**base, "converted_amount": 3_000_000_000, "remaining_amount": 5_158_000_000}, headers=h)
    assert r.status_code == 422
    # 실패한 시도는 버전을 만들지 않음
    versions = client.get(f"/api/v1/loans/{loan_id}/schedule/versions", headers=h).json()
    assert len(versions) == 1


def test_full_refinancing_matches_attachment_7():
    if not Path(EXCEL7).exists():
        pytest.skip("첨부7 파일이 없어 스킵")
    h = _headers()
    S = _read_sheet(EXCEL7, "2009-03-09_주택(대환)", last=300)
    ev_idx = next(i for i, x in enumerate(S) if x["row"] == 88)  # 2024-06-07 전환 행
    ledger, inst = [], 0
    for i, x in enumerate(S):
        memo = i + 1 < len(S) and S[i + 1]["due"] == x["due"]
        if not memo:
            inst += 1
        ledger.append((x["due"], x["start"], x["end"], x["days"], None if memo else inst,
                       x["bal"], 0 if memo else x["J"], x["K"] or x["L"], memo))
    loan_id, funds = _loan_with_ledger(h, ledger, assigned=7_500_000_000, assigned_date="2009-03-09", rate=4.643)

    r = client.post(f"/api/v1/loans/{loan_id}/events/refinancing", json={
        "refi_type": "FULL", "event_date": "2024-06-07", "new_fund_id": funds["주택도시기금"]["id"],
        "new_annual_rate": "2.34500", "rate_table_name": "2024년 2분기 대환"}, headers=h)
    assert r.status_code == 201, r.text
    v2 = r.json()
    assert v2["status"] == "NORMAL" and v2["validation_errors"] is None
    assert len(v2["rows"]) == len(S)
    assert _compare(v2["rows"], S, ev_idx + 1) == []   # 대환 이후 19개 행 전부 일치
    assert v2["rows"][-1]["balance_after"] == 0

    # 같은 원장을 이어가므로 새 차입금이 생기지 않고, 재원/금리 이력만 바뀐다
    loan = client.get(f"/api/v1/loans/{loan_id}", headers=h).json()
    assert loan["fund_id"] == funds["주택도시기금"]["id"]
    assert [x["annual_rate"] for x in loan["rate_history"]][-1] in ("2.34500", 2.345, "2.345")
    versions = client.get(f"/api/v1/loans/{loan_id}/schedule/versions", headers=h).json()
    assert [v["version_no"] for v in versions] == [1, 2]
    assert versions[1]["generated_reason"] == "대환대출(전액)"


def test_refinancing_rejects_equal_payment_ledger():
    h = _headers()
    funds, cyc, mth = _refs(h)
    run = uuid.uuid4().hex[:6]
    spc = client.post("/api/v1/spc", json={"spc_name": f"S-{run}"}, headers=h).json()["id"]
    dorm = client.post("/api/v1/dormitories", json={"dormitory_name": f"D-{run}", "school_name": "t", "spc_id": spc}, headers=h).json()["id"]
    r = client.post("/api/v1/loans", json={
        "dormitory_id": dorm, "fund_id": funds["주택도시기금"]["id"], "loan_code": f"RFP-{run}",
        "assigned_amount": 100_000_000, "assigned_date": "2023-02-28", "rate_type": "FIXED",
        "repay_cycle_code_id": cyc["SEMIANNUAL"]["id"], "repay_method_code_id": mth["EQUAL_PAYMENT"]["id"],
        "principal_repay_start_date": "2023-08-31", "repay_period_months": 60,
        "initial_rate": {"effective_start_date": "2023-02-28", "annual_rate": "2.0"}}, headers=h)
    loan_id = r.json()["id"]
    client.post(f"/api/v1/loans/{loan_id}/schedule/generate", headers=h)
    r = client.post(f"/api/v1/loans/{loan_id}/events/refinancing", json={
        "refi_type": "FULL", "event_date": "2024-08-31", "new_fund_id": funds["공공자금관리기금"]["id"],
        "new_annual_rate": "3.0"}, headers=h)
    assert r.status_code == 422
